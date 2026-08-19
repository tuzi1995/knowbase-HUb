#!/usr/bin/env python3
"""Validate and import the authoritative V1 XLSX snapshot atomically.

Default mode is read-only and writes a comparison report. ``--apply`` performs
one PostgreSQL transaction after creating a JSON backup of the affected state.
Rows missing from the XLSX are reported but are never deleted automatically.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

import psycopg2
from psycopg2.extras import Json, RealDictCursor
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "supabase_config_local.json"
DEFAULT_XLSX = Path("/Users/guoying/Downloads/最新KB（合并机型）导出_20260818_100044.xlsx")
REPORT_DIR = ROOT / "DevTools" / "reports"
BACKUP_ROOT = Path(os.environ.get("KMATRIX_BACKUP_ROOT", "/Volumes/ORICO/database/knowbasehub-backups")).expanduser()
BACKUP_DIR = BACKUP_ROOT / "v1_import_backups"
SHANGHAI = timezone(timedelta(hours=8))

IMPORT_FIELDS = [
    "question", "question_type", "answer", "answer_type", "if_bm25",
    "similar_questions", "error_list", "keyword_list", "image_urls",
    "video_urls", "file_urls", "link_type", "link_url", "update_time",
    "product_name",
]
LIST_FIELDS = {"similar_questions", "error_list", "keyword_list", "image_urls", "video_urls", "file_urls"}


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize_model(value: Any) -> str:
    return re.sub(r"\s+", " ", text(value)).casefold()


def split_models(value: Any) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[,，\n]+", text(value)):
        key = normalize_model(part)
        if key and key not in seen:
            seen.add(key)
            out.append(text(part))
    return out


def split_list(value: Any) -> list[str] | None:
    parts = [text(item) for item in re.split(r"[,，\n]+", text(value)) if text(item)]
    return parts or None


def parse_datetime(value: Any) -> datetime | None:
    if value is None or text(value) == "":
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        raw = text(value).replace("Z", "+00:00")
        try:
            dt = datetime.fromisoformat(raw)
        except ValueError:
            return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=SHANGHAI)
    return dt.astimezone(timezone.utc)


def iso(value: Any) -> str | None:
    dt = parse_datetime(value)
    return dt.isoformat() if dt else None


def canonical_json(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (list, dict, bool, int, float)):
        return value
    raw = text(value)
    if not raw:
        return None
    try:
        return json.loads(raw)
    except Exception:
        return split_list(raw)


def comparable(value: Any, field: str) -> Any:
    if field == "update_time":
        return iso(value)
    if field == "if_bm25":
        return text(value).lower() in {"1", "true", "yes", "是"} if not isinstance(value, bool) else value
    if field == "product_name":
        source = value if isinstance(value, list) else split_models(value)
        return sorted({normalize_model(x) for x in source if normalize_model(x)})
    if field in LIST_FIELDS:
        data = canonical_json(value)
        if isinstance(data, list):
            return [text(x) for x in data if text(x)]
        return [] if data is None else [text(data)]
    return text(value)


def load_snapshot(path: Path) -> tuple[dict[str, dict[str, Any]], str, list[str]]:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != [workbook.sheetnames[0]]:
        raise ValueError(f"XLSX 必须只有一个数据工作表，实际为: {workbook.sheetnames}")
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows, ())]
    required = {"update_time", "问题编号", "问题", "问题类型", "答案", "答案类型", "bm25", "机型"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"缺少必需列: {missing}")
    index = {name: pos for pos, name in enumerate(headers)}
    records: dict[str, dict[str, Any]] = {}
    duplicate_ids: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        values = list(row)
        wiki_id = text(values[index["问题编号"]])
        if not wiki_id:
            raise ValueError(f"第 {row_number} 行缺少问题编号")
        if wiki_id in records:
            duplicate_ids.append(wiki_id)
            continue
        record: dict[str, Any] = {
            "question_wiki_id": wiki_id,
            "update_time": parse_datetime(values[index["update_time"]]),
            "question": text(values[index["问题"]]),
            "question_type": text(values[index["问题类型"]]),
            "answer": text(values[index["答案"]]),
            "answer_type": text(values[index["答案类型"]]),
            "if_bm25": text(values[index["bm25"]]).lower() in {"1", "true", "yes", "是"},
            "product_name": split_models(values[index["机型"]]),
        }
        for field, header in {
            "error_list": "错误列表", "keyword_list": "关键词", "similar_questions": "相似提问",
            "image_urls": "图片链接", "video_urls": "视频链接", "file_urls": "文件链接",
            "link_type": "跳转链接类型", "link_url": "跳转链接（url/key）",
        }.items():
            value = values[index[header]] if header in index else None
            record[field] = split_list(value) if field in LIST_FIELDS else (text(value) or None)
        records[wiki_id] = record
    if duplicate_ids:
        raise ValueError(f"发现重复问题编号 {len(duplicate_ids)} 个，例如: {duplicate_ids[:5]}")
    return records, digest, headers


def connect():
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))["local_db"]
    return psycopg2.connect(**config)


def fetch_db(conn, ids: set[str] | None = None):
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        if ids is None:
            cur.execute("SELECT * FROM public.knowledge_base_v1 ORDER BY question_wiki_id")
        else:
            cur.execute("SELECT * FROM public.knowledge_base_v1 WHERE question_wiki_id = ANY(%s)", (list(ids),))
        v1 = {str(row["question_wiki_id"]): dict(row) for row in cur.fetchall()}
        cur.execute("SELECT question_wiki_id, product_name, is_configured FROM public.product_matrix WHERE question_wiki_id = ANY(%s)", (list(ids or v1.keys()),))
        matrix: dict[str, dict[str, bool]] = defaultdict(dict)
        for row in cur.fetchall():
            matrix[str(row["question_wiki_id"])][normalize_model(row["product_name"])] = bool(row["is_configured"])
        cur.execute("SELECT product_name FROM public.matrix_column ORDER BY sort_order, id")
        columns = [str(row["product_name"]) for row in cur.fetchall()]
    return v1, matrix, columns


def canonicalize_models(records: dict[str, dict[str, Any]], columns: list[str]) -> dict[str, str]:
    by_key = {normalize_model(name): name for name in columns}
    compact_map: dict[str, list[str]] = defaultdict(list)
    for key, name in by_key.items():
        compact_map[re.sub(r"\s+", "", key)].append(name)
    missing: dict[str, str] = {}
    for record in records.values():
        canonical: list[str] = []
        for raw in record["product_name"]:
            key = normalize_model(raw)
            if key not in by_key:
                compact_matches = compact_map.get(re.sub(r"\s+", "", key), [])
                if len(compact_matches) == 1:
                    by_key[key] = compact_matches[0]
            if key not in by_key:
                missing[key] = raw
            else:
                canonical.append(by_key[key])
        record["product_name"] = canonical
    if missing:
        raise ValueError(f"文件包含型号字典中不存在的型号: {sorted(missing.values())}")
    return by_key


def build_diff(records, db, matrix):
    rows = []
    fields_counter: Counter[str] = Counter()
    file_ids, db_ids = set(records), set(db)
    for wiki_id in sorted(file_ids | db_ids):
        incoming = records.get(wiki_id)
        current = db.get(wiki_id)
        if incoming is None:
            rows.append({"question_wiki_id": wiki_id, "action": "keep_db_only", "changed_fields": "", "note": "文件未包含，按方案不自动删除"})
            continue
        if current is None:
            changed = IMPORT_FIELDS[:]
            rows.append({"question_wiki_id": wiki_id, "action": "insert", "changed_fields": ",".join(changed), "note": "文件新增 V1 条目"})
            fields_counter.update(changed)
            continue
        changed = [field for field in IMPORT_FIELDS if comparable(incoming.get(field), field) != comparable(current.get(field), field)]
        if changed:
            fields_counter.update(changed)
            rows.append({"question_wiki_id": wiki_id, "action": "update", "changed_fields": ",".join(changed), "note": "以 XLSX 当前值覆盖对应 V1 字段"})
        elif sorted(key for key, configured in matrix.get(wiki_id, {}).items() if configured) != sorted({normalize_model(model) for model in incoming["product_name"]}):
            rows.append({"question_wiki_id": wiki_id, "action": "matrix_sync", "changed_fields": "product_name", "note": "V1 已一致，仅同步落后的矩阵投影"})
    return rows, fields_counter


def write_report(path: Path, summary: dict[str, Any], rows: list[dict[str, Any]]):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"summary": summary, "rows": rows}, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    csv_path = path.with_suffix(".csv")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["question_wiki_id", "action", "changed_fields", "note"])
        writer.writeheader()
        writer.writerows(rows)


def serializable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: serializable(item) for key, item in value.items()}
    if isinstance(value, list):
        return [serializable(item) for item in value]
    return value


def apply_import(conn, records, db, matrix, operation_id, xlsx_path, digest, diff_rows):
    changed_ids = {row["question_wiki_id"] for row in diff_rows if row["action"] in {"insert", "update", "matrix_sync"}}
    backup_dir = BACKUP_DIR / operation_id
    backup_dir.mkdir(parents=True, exist_ok=True)
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM public.product_matrix WHERE question_wiki_id = ANY(%s)", (list(changed_ids),))
        matrix_backup = [dict(row) for row in cur.fetchall()]
        cur.execute("SELECT * FROM public.matrix_column ORDER BY sort_order, id")
        columns_backup = [dict(row) for row in cur.fetchall()]
    (backup_dir / "knowledge_base_v1_before.json").write_text(json.dumps([serializable(row) for row in db.values()], ensure_ascii=False, indent=2), encoding="utf-8")
    (backup_dir / "product_matrix_before.json").write_text(json.dumps([serializable(row) for row in matrix_backup], ensure_ascii=False, indent=2), encoding="utf-8")
    (backup_dir / "matrix_column_before.json").write_text(json.dumps([serializable(row) for row in columns_backup], ensure_ascii=False, indent=2), encoding="utf-8")

    now = datetime.now(timezone.utc)
    audit_count = 0
    matrix_changed_count = 0
    with conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT product_name FROM public.matrix_column")
            column_keys = {normalize_model(row["product_name"]) for row in cur.fetchall()}
            for record in records.values():
                for model in record["product_name"]:
                    if normalize_model(model) not in column_keys:
                        raise RuntimeError(f"写入前发现未登记型号: {model}")

            for row in diff_rows:
                if row["action"] not in {"insert", "update"}:
                    continue
                wiki_id = row["question_wiki_id"]
                incoming = records[wiki_id]
                current = db.get(wiki_id)
                before_version = int((current or {}).get("content_version") or 0) if current else 0
                changed_fields = [field for field in IMPORT_FIELDS if comparable(incoming.get(field), field) != comparable((current or {}).get(field), field)] if current else IMPORT_FIELDS[:]
                committed_version = before_version + 1 if current else 1
                product_text = ",".join(incoming["product_name"])
                values = {
                    "question_wiki_id": wiki_id,
                    "question_type": incoming["question_type"], "question": incoming["question"], "answer": incoming["answer"],
                    "answer_type": incoming["answer_type"], "if_bm25": incoming["if_bm25"],
                    "similar_questions": incoming["similar_questions"], "error_list": incoming["error_list"], "keyword_list": incoming["keyword_list"],
                    "image_urls": incoming["image_urls"], "video_urls": incoming["video_urls"], "file_urls": incoming["file_urls"],
                    "link_type": incoming["link_type"], "link_url": incoming["link_url"], "update_time": incoming["update_time"],
                    "product_name": product_text,
                    "product_category_name": (current or {}).get("product_category_name"),
                    "combinatorial_problem": (current or {}).get("combinatorial_problem"),
                    "review_status": "unadjusted",
                    "content_version": committed_version,
                    "content_lifecycle_status": (current or {}).get("content_lifecycle_status") or "normal",
                }
                cur.execute(
                    """
                    INSERT INTO public.knowledge_base_v1 (
                      question_wiki_id, question_type, question, answer, answer_type, if_bm25,
                      similar_questions, error_list, keyword_list, image_urls, video_urls, file_urls,
                      link_type, link_url, update_time, product_category_name, product_name,
                      combinatorial_problem, review_status, content_version, content_lifecycle_status
                    ) VALUES (%(question_wiki_id)s, %(question_type)s, %(question)s, %(answer)s, %(answer_type)s, %(if_bm25)s,
                      %(similar_questions)s, %(error_list)s, %(keyword_list)s, %(image_urls)s, %(video_urls)s, %(file_urls)s,
                      %(link_type)s, %(link_url)s, %(update_time)s, %(product_category_name)s, %(product_name)s,
                      %(combinatorial_problem)s, %(review_status)s, %(content_version)s, %(content_lifecycle_status)s)
                    ON CONFLICT (question_wiki_id) DO UPDATE SET
                      question_type=EXCLUDED.question_type, question=EXCLUDED.question, answer=EXCLUDED.answer,
                      answer_type=EXCLUDED.answer_type, if_bm25=EXCLUDED.if_bm25, similar_questions=EXCLUDED.similar_questions,
                      error_list=EXCLUDED.error_list, keyword_list=EXCLUDED.keyword_list, image_urls=EXCLUDED.image_urls,
                      video_urls=EXCLUDED.video_urls, file_urls=EXCLUDED.file_urls, link_type=EXCLUDED.link_type,
                      link_url=EXCLUDED.link_url, update_time=EXCLUDED.update_time, product_name=EXCLUDED.product_name,
                      review_status=EXCLUDED.review_status, content_version=EXCLUDED.content_version,
                      content_lifecycle_status=EXCLUDED.content_lifecycle_status
                    """,
                    {**values,
                     "similar_questions": Json(values["similar_questions"]) if values["similar_questions"] is not None else None,
                     "error_list": Json(values["error_list"]) if values["error_list"] is not None else None,
                     "keyword_list": Json(values["keyword_list"]) if values["keyword_list"] is not None else None,
                     "image_urls": Json(values["image_urls"]) if values["image_urls"] is not None else None,
                     "video_urls": Json(values["video_urls"]) if values["video_urls"] is not None else None,
                     "file_urls": Json(values["file_urls"]) if values["file_urls"] is not None else None},
                )

                before_snapshot = {field: serializable((current or {}).get(field)) for field in IMPORT_FIELDS} if current else {}
                after_snapshot = {field: serializable(incoming.get(field)) for field in IMPORT_FIELDS}
                cur.execute(
                    """
                    INSERT INTO public.knowledge_base_modifications (
                      kb_id, question_wiki_id, product_name, product_category_name, question, answer,
                      similar_questions, keyword_list, image_urls, video_urls, file_urls, link_type, link_url,
                      change_meta, modification_time, change_type, modifier, question_type, answer_type,
                      error_list, if_bm25, source_module, operation_id, source_code, base_version,
                      committed_version, submitted_version, conflict_resolution
                    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'edit','system_import',%s,%s,%s,%s,'全量导入',%s,'full_import',%s,%s,%s,'normal')
                    -- The idempotency index is partial; an unqualified
                    -- conflict clause lets PostgreSQL use it safely.
                    ON CONFLICT DO NOTHING
                    """,
                    (wiki_id, wiki_id, product_text, values["product_category_name"], incoming["question"], incoming["answer"],
                     Json(incoming["similar_questions"]) if incoming["similar_questions"] is not None else None,
                     Json(incoming["keyword_list"]) if incoming["keyword_list"] is not None else None,
                     Json(incoming["image_urls"]) if incoming["image_urls"] is not None else None,
                     Json(incoming["video_urls"]) if incoming["video_urls"] is not None else None,
                     Json(incoming["file_urls"]) if incoming["file_urls"] is not None else None,
                     incoming["link_type"], incoming["link_url"],
                     Json({"source": "全量导入", "source_code": "full_import", "operation_id": operation_id,
                           "input_file": str(xlsx_path), "input_sha256": digest,
                           "changed_fields": changed_fields, "before": before_snapshot, "after": after_snapshot}),
                     now, incoming["question_type"], incoming["answer_type"],
                     Json(incoming["error_list"]) if incoming["error_list"] is not None else None,
                     incoming["if_bm25"], operation_id, before_version or None, committed_version, committed_version),
                )
                audit_count += cur.rowcount

                desired = {normalize_model(model): model for model in incoming["product_name"]}
                cur.execute("SELECT product_name,is_configured FROM public.product_matrix WHERE question_wiki_id=%s", (wiki_id,))
                current_matrix = {
                    normalize_model(row["product_name"]): {
                        "is_configured": bool(row["is_configured"]),
                        "product_name": row["product_name"],
                    }
                    for row in cur.fetchall()
                }
                for key, model in desired.items():
                    cur.execute(
                        """
                        INSERT INTO public.product_matrix
                          (question_wiki_id, product_name, is_configured, manual_edit, edit_source, last_synced_at,
                           question_content, answer_content, update_time, product_category)
                        VALUES (%s,%s,true,true,'full_import',now(),%s,%s,%s,%s)
                        ON CONFLICT (question_wiki_id, product_name) DO UPDATE SET
                          is_configured=true, manual_edit=true, edit_source='full_import', last_synced_at=now(),
                          question_content=EXCLUDED.question_content, answer_content=EXCLUDED.answer_content,
                          update_time=EXCLUDED.update_time, product_category=EXCLUDED.product_category
                        """,
                        (wiki_id, model, incoming["question"], incoming["answer"], incoming["update_time"].isoformat() if incoming["update_time"] else None, values["product_category_name"]),
                    )
                stale = [key for key, item in current_matrix.items() if item["is_configured"] and key not in desired]
                for key in stale:
                    existing_name = current_matrix[key]["product_name"]
                    cur.execute("UPDATE public.product_matrix SET is_configured=false, manual_edit=true, edit_source='full_import', last_synced_at=now() WHERE question_wiki_id=%s AND lower(regexp_replace(product_name, '\\s+', ' ', 'g'))=lower(regexp_replace(%s, '\\s+', ' ', 'g'))", (wiki_id, existing_name))
                if stale or set(current_matrix) != set(desired):
                    matrix_changed_count += 1

            # Some rows already matched the XLSX in V1 while their legacy
            # matrix projection was stale. Sync those rows without creating a
            # second V1 audit event or incrementing content_version.
            for row in diff_rows:
                if row["action"] != "matrix_sync":
                    continue
                wiki_id = row["question_wiki_id"]
                incoming = records[wiki_id]
                desired = {normalize_model(model): model for model in incoming["product_name"]}
                cur.execute("SELECT product_name,is_configured FROM public.product_matrix WHERE question_wiki_id=%s", (wiki_id,))
                current_matrix = {
                    normalize_model(item["product_name"]): {
                        "is_configured": bool(item["is_configured"]),
                        "product_name": item["product_name"],
                    }
                    for item in cur.fetchall()
                }
                for model in desired.values():
                    cur.execute(
                        """
                        INSERT INTO public.product_matrix
                          (question_wiki_id, product_name, is_configured, manual_edit, edit_source, last_synced_at,
                           question_content, answer_content, update_time, product_category)
                        VALUES (%s,%s,true,true,'full_import',now(),%s,%s,%s,%s)
                        ON CONFLICT (question_wiki_id, product_name) DO UPDATE SET
                          is_configured=true, manual_edit=true, edit_source='full_import', last_synced_at=now(),
                          question_content=EXCLUDED.question_content, answer_content=EXCLUDED.answer_content,
                          update_time=EXCLUDED.update_time
                        """,
                        (wiki_id, model, incoming["question"], incoming["answer"], incoming["update_time"].isoformat() if incoming["update_time"] else None, None),
                    )
                for key, item in current_matrix.items():
                    if item["is_configured"] and key not in desired:
                        cur.execute("UPDATE public.product_matrix SET is_configured=false, manual_edit=true, edit_source='full_import', last_synced_at=now() WHERE question_wiki_id=%s AND product_name=%s", (wiki_id, item["product_name"]))
                matrix_changed_count += 1

    return {"operation_id": operation_id, "backup_dir": str(backup_dir), "audit_rows_inserted": audit_count, "matrix_rows_changed": matrix_changed_count}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--apply", action="store_true", help="apply after validation; default is read-only")
    args = parser.parse_args()
    records, digest, headers = load_snapshot(args.xlsx)
    conn = connect()
    try:
        db, matrix, columns = fetch_db(conn)
        canonicalize_models(records, columns)
        diff_rows, field_counts = build_diff(records, db, matrix)
        action_counts = Counter(row["action"] for row in diff_rows)
        summary: dict[str, Any] = {
            "xlsx": str(args.xlsx), "sha256": digest, "headers": headers,
            "incoming_rows": len(records), "current_v1_rows": len(db),
            "insert_rows": action_counts.get("insert", 0), "update_rows": action_counts.get("update", 0),
            "db_only_rows_kept": action_counts.get("keep_db_only", 0), "changed_field_counts": dict(field_counts),
            "operation_id": f"full_import_{digest[:16]}", "applied": False,
        }
        report_path = REPORT_DIR / f"latest_v1_xlsx_import_{digest[:16]}.json"
        write_report(report_path, summary, diff_rows)
        print(json.dumps({"summary": summary, "report": str(report_path), "csv": str(report_path.with_suffix('.csv'))}, ensure_ascii=False, indent=2))
        if args.apply:
            result = apply_import(conn, records, db, matrix, summary["operation_id"], args.xlsx, digest, diff_rows)
            summary.update(result, applied=True)
            write_report(report_path, summary, diff_rows)
            print(json.dumps({"applied": True, **result}, ensure_ascii=False, indent=2))
        return 0
    finally:
        conn.close()


if __name__ == "__main__":
    raise SystemExit(main())
