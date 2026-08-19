import io
import os
import sqlite3
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook


_TEMP_DIR = None
if not os.environ.get("KMATRIX_SQLITE_PATH"):
    _TEMP_DIR = tempfile.TemporaryDirectory()
    os.environ["KMATRIX_SQLITE_PATH"] = os.path.join(_TEMP_DIR.name, "clone-review.db")

import server  # noqa: E402


class MatrixCloneReviewApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        server.app.config.update(TESTING=True)
        server.init_db()
        cls._database_client_patcher = patch.object(server, "get_supabase_client", return_value=None)
        cls._database_client_patcher.start()

    @classmethod
    def tearDownClass(cls):
        cls._database_client_patcher.stop()
        with server.app.app_context():
            server.db.session.remove()

    def setUp(self):
        with server.app.app_context():
            server.Button.query.delete()
            server.MatrixSubmitOperation.query.delete()
            server.MatrixCloneReviewItem.query.delete()
            server.MatrixCloneDifferenceRule.query.delete()
            server.MatrixCloneReviewBatch.query.delete()
            server.ProductMatrix.query.delete()
            server.MatrixColumn.query.delete()
            server.db.session.add_all([
                server.MatrixColumn(product_name="P20", sort_order=1),
                server.MatrixColumn(product_name="P30 Pro", sort_order=2),
                server.ProductMatrix(
                    question_wiki_id="KB-OLD",
                    product_name="P20",
                    is_configured=True,
                    question_content="最大吸力是多少？",
                    answer_content="最大吸力为 22000Pa。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-NEW",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="最大吸力是多少？",
                    answer_content="最大吸力为 40000Pa。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-WRONG",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="电机额定功率是多少？",
                    answer_content="额定功率为 40000Pa。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-DISABLED",
                    product_name="P20",
                    is_configured=False,
                    question_content="已停用的知识不应计入来源统计。",
                    answer_content="停用。",
                    product_category="扫地机",
                ),
            ])
            server.db.session.commit()
        self.client = server.app.test_client()
        login = self.client.post("/login", json={"username": "admin", "password": "123456"})
        self.assertEqual(login.status_code, 200)

    def _clone(self, operation_id="clone-op-1"):
        refresh = {
            "snapshot_id": "snapshot-test-latest",
            "source_version": "v-test-latest",
            "comparison_run_id": "run-test-latest",
            "scan_run_id": "scan-test-latest",
            "result_counts": {},
        }
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "load_ai_config", return_value={}
        ), patch.object(server, "_refresh_clone_parameter_comparison", return_value=refresh):
            return self.client.post("/api/matrix/clone_config", json={
                "operation_id": operation_id,
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True,
                    "auto_scan": True,
                    "detection_methods": ["manual_difference_ai"],
                    "manual_difference_text": "旧型号最大吸力 22000Pa，新品 40000Pa",
                },
            })

    @staticmethod
    def _tag_client(mappings):
        client = Mock()

        def select_all(table, **kwargs):
            if table == "kb_tags":
                return [
                    {"id": "tag-skip", "name": "不克隆"},
                    {"id": "tag-risk", "name": "高风险"},
                ]
            if table == "kb_item_tags":
                if kwargs.get("filters", {}).get("library_type") != "eq.current":
                    raise AssertionError("clone tag exclusions must use the current knowledge library")
                return list(mappings)
            raise AssertionError(f"unexpected table: {table}")

        client.select_all.side_effect = select_all
        return client

    def test_clone_preview_excludes_any_selected_tag_once(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-SECOND",
                product_name="P20",
                is_configured=True,
                question_content="第二条知识",
                answer_content="保留",
                product_category="扫地机",
            ))
            server.db.session.commit()

        tag_client = self._tag_client([
            {"question_wiki_id": "KB-OLD", "tag_id": "tag-skip"},
            {"question_wiki_id": "KB-OLD", "tag_id": "tag-risk"},
        ])
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "get_supabase_client", return_value=tag_client
        ):
            response = self.client.post("/api/matrix/clone_config/preview", json={
                "mode": "model",
                "source": "P20",
                "scope": {"mode": "all", "exclude_tag_names": ["不克隆", "高风险"]},
            })

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        body = response.get_json()
        self.assertEqual(body["source_count"], 1)
        self.assertEqual(body["excluded_count"], 1)
        self.assertEqual(body["exclude_tag_names"], ["不克隆", "高风险"])

    def test_clone_does_not_create_target_relation_for_excluded_tagged_kb(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-TAGGED",
                product_name="P20",
                is_configured=True,
                question_content="不应克隆",
                answer_content="排除",
                product_category="扫地机",
            ))
            server.db.session.commit()

        tag_client = self._tag_client([
            {"question_wiki_id": "KB-TAGGED", "tag_id": "tag-skip"},
        ])
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "get_supabase_client", return_value=tag_client
        ):
            response = self.client.post("/api/matrix/clone_config", json={
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all", "exclude_tag_names": ["不克隆"]},
            })

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        body = response.get_json()
        self.assertEqual(body["source_count"], 1)
        self.assertEqual(body["excluded_count"], 1)
        with server.app.app_context():
            self.assertIsNotNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P30 Pro", is_configured=True,
            ).first())
            self.assertIsNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-TAGGED", product_name="P30 Pro",
            ).first())

    def test_clone_rejects_deleted_exclusion_tag_before_creating_target(self):
        tag_client = self._tag_client([])
        with patch.object(server, "get_supabase_client", return_value=tag_client):
            response = self.client.post("/api/matrix/clone_config", json={
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all", "exclude_tag_names": ["已删除标签"]},
            })

        self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["code"], "clone_exclude_tags_invalid")
        with server.app.app_context():
            self.assertIsNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P30 Pro",
            ).first())

    def test_parameter_review_refreshes_latest_data_and_binds_the_new_run(self):
        refresh = {
            "snapshot_id": "snapshot-latest",
            "source_version": "v-latest",
            "comparison_run_id": "run-latest",
            "scan_run_id": "scan-latest",
            "result_counts": {"finding_count": 1},
        }
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "_refresh_clone_parameter_comparison", return_value=refresh
        ) as refresh_call, patch.object(
            server, "_clone_review_parameter_difference_rules", return_value=([], None)
        ) as rules_call:
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-parameter-latest",
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True,
                    "auto_scan": True,
                    "detection_methods": ["parameter_check"],
                },
            })

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        body = response.get_json()
        self.assertEqual(body["parameter_refresh"]["comparison_run_id"], "run-latest")
        refresh_call.assert_called_once_with("扫地机")
        rules_call.assert_called_once_with("P20", "P30 Pro")
        detail = self.client.get(f"/api/matrix/clone-review/{body['review_batch_id']}").get_json()
        self.assertEqual(detail["parameter_refresh"]["snapshot_id"], "snapshot-latest")

    def test_parameter_review_uses_category_shared_by_source_and_target_models(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-FLOOR-WASHER",
                product_name="A30 Pro Ultra",
                is_configured=True,
                question_content="洗地机功能是什么？",
                answer_content="源型号说明。",
                product_category="洗地机",
            ))
            server.db.session.commit()

        refresh = {
            "snapshot_id": "snapshot-floor-washer",
            "source_version": "v-floor-washer",
            "comparison_run_id": "run-floor-washer",
            "scan_run_id": "scan-floor-washer",
            "result_counts": {},
        }
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "_refresh_clone_parameter_comparison", return_value=refresh
        ) as refresh_call, patch.object(
            server, "_clone_review_parameter_difference_rules", return_value=([], None)
        ):
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-floor-washer-parameters",
                "mode": "model",
                "source": "A30 Pro Ultra",
                "targets": ["A30 Pro Ultra 3.0"],
                "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True,
                    "auto_scan": True,
                    "detection_methods": ["parameter_check"],
                },
            })

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        refresh_call.assert_called_once_with("洗地机")

    def test_parameter_category_requires_one_shared_catalog_category(self):
        with patch.object(server, "parse_product_catalog", return_value={
            "扫地机": ["共享源", "共享目标"],
            "洗地机": ["共享源", "共享目标"],
        }):
            with self.assertRaisesRegex(ValueError, "同一唯一产品品类"):
                server._clone_review_parameter_category("共享源", "共享目标")

    def test_parameter_refresh_uses_explicit_worker_actor_without_current_user(self):
        snapshot = {"snapshot_id": "snapshot-1", "source_version": "v1", "content_hash": "hash"}
        comparison = {"run_id": "run-1", "scan_result": {"run_id": "scan-1"}, "result_counts": {}}
        connection = Mock()
        with patch.object(server, "current_user", None), patch.object(
            server, "fetch_catalog_snapshot", return_value=snapshot
        ), patch.object(server, "connect_main_database", return_value=connection), patch.object(
            server, "store_snapshot"
        ) as store_snapshot, patch.object(
            server, "refresh_model_bindings", return_value={}
        ) as refresh_bindings, patch.object(
            server, "run_current_knowledge_parameter_comparison", return_value=comparison
        ) as run_comparison:
            result = server._refresh_clone_parameter_comparison("扫地机", actor="admin")

        self.assertEqual(result["comparison_run_id"], "run-1")
        store_snapshot.assert_called_once_with(connection, snapshot, "admin")
        self.assertEqual(refresh_bindings.call_args.kwargs["actor"], "admin")
        self.assertEqual(run_comparison.call_args.kwargs["actor"], "admin")
        connection.close.assert_called_once()

    def test_risk_summary_partitions_only_scanned_high_and_medium_items(self):
        def item(
            item_id,
            detection_status,
            risk_level,
            decision="pending",
            apply_status="not_applied",
            decided_by=None,
        ):
            return server.MatrixCloneReviewItem(
                item_id=item_id,
                batch_id="risk-summary-batch",
                question_wiki_id=item_id,
                target_model="P30 Pro",
                source_model="P20",
                detection_status=detection_status,
                risk_level=risk_level,
                decision=decision,
                apply_status=apply_status,
                decided_by=decided_by,
            )

        items = [
            item("high-pending", "scanned", "high"),
            item("medium-pending", "scanned", "medium"),
            item("high-awaiting-apply", "scanned", "high", "link_existing", "not_applied", "admin"),
            item("high-processed", "scanned", "high", "keep", "not_required"),
            item("medium-processed", "scanned", "medium", "remove_target", "applied"),
            item("default-processed", "scanned", "default", "keep", "not_required"),
            item("scan-failed", "scan_failed", "unknown"),
            item("waiting", "pending_scan", "unknown"),
        ]

        summary = server._clone_review_risk_summary(items)

        self.assertEqual(summary, {"total": 5, "pending": 3, "processed": 2})
        self.assertEqual(summary["pending"] + summary["processed"], summary["total"])

        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="risk-summary-batch",
                operation_id="risk-summary-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json='[]',
                status="reviewing",
                source_count=len(items),
                item_count=len(items),
                created_by="admin",
            )
            server.db.session.add(batch)
            server.db.session.add_all(items)
            server.db.session.commit()

        payload = self.client.get("/api/matrix/clone-review/risk-summary-batch").get_json()
        self.assertEqual(payload["risk_summary"], summary)

    def test_async_parameter_review_passes_request_actor_to_worker(self):
        with patch.object(server, "_matrix_clone_review_async_enabled", return_value=True), patch.object(
            server, "threading"
        ) as threading_mock:
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-async-actor",
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True,
                    "auto_scan": True,
                    "detection_methods": ["parameter_check"],
                },
            })

        self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
        args = threading_mock.Thread.call_args.kwargs["args"]
        self.assertEqual(args[1:], ({"mode": "all"}, True, "admin"))
        threading_mock.Thread.return_value.start.assert_called_once()

    def test_ai_assists_unseen_wording_without_overriding_deterministic_rules(self):
        ai_text = "P30 Pro：支持激光边刷 P20 Ultra活水版：普通边刷"
        ai_response = {
            "rules": [
                {
                    "feature_id": "max_suction",
                    "feature_name": "最大吸力",
                    "source_value": "22000Pa",
                    "target_value": "50000Pa",
                    "evidence_quote": "旧型号最大吸力 22000Pa，新品 50000Pa",
                },
                {
                    "feature_id": "edge_brush_type",
                    "feature_name": "边刷类型",
                    "source_value": "普通边刷",
                    "target_value": "激光边刷",
                    "evidence_quote": ai_text,
                    "confidence": 0.93,
                },
            ]
        }
        with patch.dict(os.environ, {"MATRIX_CLONE_REVIEW_AI_SYNC": "1"}), patch.object(
            server, "load_ai_config", return_value={"api_key": "configured", "base_url": "http://ai.test", "model": "test-model"}
        ), patch.object(server, "_ai_call_llm", return_value=__import__("json").dumps(ai_response, ensure_ascii=False)) as call:
            rules, status = server._clone_review_ai_rules(
                "旧型号最大吸力 22000Pa，新品 40000Pa\n" + ai_text,
                "P30 Pro",
            )

        call.assert_called_once()
        self.assertEqual(status, "completed_ai")
        by_feature = {rule["feature_id"]: rule for rule in rules}
        self.assertEqual(by_feature["max_suction"]["target_value"], "40000Pa")
        self.assertEqual(by_feature["edge_brush_type"]["source_value"], "普通边刷")
        self.assertIn("configured_ai", by_feature["edge_brush_type"]["evidence_sources"])

    def test_ai_status_distinguishes_no_change_fallback_and_deterministic_only(self):
        manual_text = "旧型号最大吸力 22000Pa，新品 40000Pa"
        config = {"api_key": "configured", "base_url": "http://ai.test", "model": "test-model"}
        with patch.dict(os.environ, {"MATRIX_CLONE_REVIEW_AI_SYNC": "1"}), patch.object(
            server, "load_ai_config", return_value=config
        ), patch.object(server, "_ai_call_llm", return_value='{"rules": []}'):
            _, no_change_status = server._clone_review_ai_rules(manual_text, "P30 Pro")

        with patch.dict(os.environ, {"MATRIX_CLONE_REVIEW_AI_SYNC": "1"}), patch.object(
            server, "load_ai_config", return_value=config
        ), patch.object(server, "_ai_call_llm", side_effect=RuntimeError("AI unavailable")):
            fallback_rules, fallback_status = server._clone_review_ai_rules(manual_text, "P30 Pro")

        with patch.dict(os.environ, {"MATRIX_CLONE_REVIEW_AI_SYNC": "0"}), patch.object(
            server, "_ai_call_llm"
        ) as disabled_call:
            _, deterministic_status = server._clone_review_ai_rules(manual_text, "P30 Pro")

        self.assertEqual(no_change_status, "completed_ai_no_change")
        self.assertEqual(fallback_status, "completed_ai_fallback")
        self.assertEqual(fallback_rules[0]["target_value"], "40000Pa")
        self.assertEqual(deterministic_status, "completed_deterministic")
        disabled_call.assert_not_called()

    def test_ai_failure_without_a_deterministic_rule_is_parse_failed(self):
        with patch.dict(os.environ, {"MATRIX_CLONE_REVIEW_AI_SYNC": "1"}), patch.object(
            server, "load_ai_config", return_value={"api_key": "configured", "base_url": "http://ai.test", "model": "test-model"}
        ), patch.object(server, "_ai_call_llm", side_effect=RuntimeError("AI unavailable")):
            rules, status = server._clone_review_ai_rules("新品和旧型号存在差异", "P30 Pro")

        self.assertEqual(rules, [])
        self.assertEqual(status, "parse_failed")

    def test_clone_creates_idempotent_review_and_blocks_early_submit(self):
        first = self._clone()
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        body = first.get_json()
        self.assertEqual(body["pending_count"], 1)
        self.assertEqual(body["review_status"], "reviewing")
        with server.app.app_context():
            self.assertIsNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P30 Pro",
            ).first())

        second = self._clone()
        self.assertEqual(second.status_code, 200)
        self.assertEqual(second.get_json()["review_batch_id"], body["review_batch_id"])

        detail = self.client.get(f"/api/matrix/clone-review/{body['review_batch_id']}").get_json()
        self.assertEqual(detail["items"][0]["risk_level"], "high")
        self.assertEqual(detail["items"][0]["reuse_candidates"][0]["wiki_id"], "KB-NEW")
        self.assertEqual(len(detail["items"][0]["reuse_candidates"]), 1)

        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            blocked = self.client.post("/api/matrix/submit_changes", json={
                "operation_id": "submit-too-early",
                "changes": [{
                    "question_wiki_id": "KB-OLD",
                    "product_name": "P30 Pro",
                    "old_is_configured": False,
                    "new_is_configured": True,
                    "edit_source": "bulk",
                }],
            })
        self.assertEqual(blocked.status_code, 409)
        self.assertEqual(blocked.get_json()["code"], "clone_review_required")

    def test_clone_rejects_reusing_an_operation_id_for_different_request(self):
        created = self._clone("clone-op-reused").get_json()
        self.assertTrue(created["success"])

        changed_request = self._clone("clone-op-reused")
        self.assertEqual(changed_request.status_code, 200)

        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "load_ai_config", return_value={}
        ), patch.object(server, "_refresh_clone_parameter_comparison", return_value={}):
            conflict = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-reused",
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "review": {"enabled": True, "detection_methods": ["manual_difference_ai"], "manual_difference_text": "不同的差异说明"},
            })

        self.assertEqual(conflict.status_code, 409, conflict.get_data(as_text=True))
        self.assertEqual(conflict.get_json()["code"], "operation_id_conflict")

    def test_submit_rejects_reusing_an_operation_id_for_different_changes(self):
        first_changes = [{
            "question_wiki_id": "KB-OLD",
            "product_name": "P30 Pro",
            "old_is_configured": False,
            "new_is_configured": True,
            "edit_source": "bulk",
        }]
        changed_request = [{
            "question_wiki_id": "KB-OLD",
            "product_name": "P30 Pro",
            "old_is_configured": False,
            "new_is_configured": True,
            "edit_source": "cell",
        }]
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-OLD",
                product_name="P30 Pro",
                is_configured=True,
                manual_edit=True,
                edit_source="bulk",
            ))
            server.db.session.commit()
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            first = self.client.post("/api/matrix/submit_changes", json={
                "operation_id": "submit-op-reused", "changes": first_changes,
            })
            conflict = self.client.post("/api/matrix/submit_changes", json={
                "operation_id": "submit-op-reused", "changes": changed_request,
            })

        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        self.assertEqual(conflict.status_code, 409, conflict.get_data(as_text=True))
        self.assertEqual(conflict.get_json()["code"], "operation_id_conflict")

    def test_source_stats_only_count_configured_matrix_rows(self):
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            response = self.client.get("/api/matrix/stats?mode=model&source=P20")

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["count"], 1)

    def test_reuse_candidates_classify_generic_semantic_numeric_expressions(self):
        with server.app.app_context():
            server.db.session.add_all([
                server.ProductMatrix(
                    question_wiki_id="KB-OBSTACLE-COMBINED",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="越障高度",
                    answer_content="单层最高 4.5cm，双层最高 4.5+4.3cm。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-HEIGHT-MM",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="机器高度是多少？",
                    answer_content="整机高度为 88mm。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-WRONG-VALUE",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="越障高度",
                    answer_content="最大越障高度为 6cm。",
                    product_category="扫地机",
                ),
            ])
            server.db.session.commit()

            obstacle_item = server.MatrixCloneReviewItem(
                question_wiki_id="KB-SOURCE-OBSTACLE",
                question="为什么门槛过不去？",
                answer="旧型号越障高度为 2cm。",
                product_category="扫地机",
            )
            obstacle_rules = [{
                "feature_id": "obstacle_crossing_height",
                "feature_name": "越障高度",
                "target_value": "8.8cm",
                "source_search_terms": ["越障高度", "门槛"],
                "target_search_terms": ["越障高度", "门槛"],
                "kb_intents": ["越障高度"],
            }]
            obstacle_candidates = server._clone_review_reuse_candidates(obstacle_item, obstacle_rules)
            obstacle_by_id = {candidate["wiki_id"]: candidate for candidate in obstacle_candidates}

            self.assertEqual(
                obstacle_by_id["KB-OBSTACLE-COMBINED"]["validation"],
                "semantic_numeric_equivalent",
            )
            self.assertIn("4.5+4.3cm", obstacle_by_id["KB-OBSTACLE-COMBINED"]["match_reason"])
            self.assertNotIn("KB-WRONG-VALUE", obstacle_by_id)

            height_item = server.MatrixCloneReviewItem(
                question_wiki_id="KB-SOURCE-HEIGHT",
                question="机身有多高？",
                answer="旧型号机身高度为 10cm。",
                product_category="扫地机",
            )
            height_rules = [{
                "feature_id": "body_height",
                "feature_name": "机身高度",
                "target_value": "8.8cm",
                "source_search_terms": ["机身高度", "机器高度"],
                "target_search_terms": ["机身高度", "机器高度"],
                "kb_intents": ["机身高度"],
            }]
            height_candidates = server._clone_review_reuse_candidates(height_item, height_rules)
            height_by_id = {candidate["wiki_id"]: candidate for candidate in height_candidates}

            self.assertEqual(height_by_id["KB-HEIGHT-MM"]["validation"], "semantic_numeric_equivalent")
            self.assertNotIn("KB-OBSTACLE-COMBINED", height_by_id)

    def test_exact_reuse_candidate_keeps_strict_validation(self):
        with server.app.app_context():
            item = server.MatrixCloneReviewItem(
                question_wiki_id="KB-OLD",
                question="最大吸力是多少？",
                answer="最大吸力为 22000Pa。",
                product_category="扫地机",
            )
            rules = [{
                "feature_id": "max_suction",
                "feature_name": "最大吸力",
                "target_value": "40000Pa",
                "source_search_terms": ["最大吸力"],
                "target_search_terms": ["最大吸力"],
                "kb_intents": ["最大吸力"],
            }]

            candidates = server._clone_review_reuse_candidates(item, rules)
            candidate = next(row for row in candidates if row["wiki_id"] == "KB-NEW")

            self.assertEqual(candidate["validation"], "exact_reusable")
            self.assertEqual(candidate["priority"], "P1")
            self.assertEqual(candidate["evidence_matches"], {
                "question": True,
                "answer": True,
                "feature_value": True,
            })

    def test_reuse_candidates_rank_three_evidence_levels_and_block_conflicts(self):
        with server.app.app_context():
            server.db.session.add_all([
                server.ProductMatrix(
                    question_wiki_id="KB-P2",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="最大吸力是多少？",
                    answer_content="40000Pa",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-P3",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="电池容量及续航",
                    answer_content="电池容量为 6400mAh。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-P3-QUESTION",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="最大吸力是多少？",
                    answer_content="请联系人工客服确认。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-CONFLICT",
                    product_name="P30-CANDIDATE-SOURCE",
                    is_configured=True,
                    question_content="最大吸力是多少？",
                    answer_content="最大吸力为 22000Pa。",
                    product_category="扫地机",
                ),
            ])
            server.db.session.commit()
            item = server.MatrixCloneReviewItem(
                question_wiki_id="KB-OLD",
                question="最大吸力是多少？",
                answer="最大吸力为 22000Pa。",
                product_category="扫地机",
            )
            rules = [{
                "feature_id": "max_suction",
                "feature_name": "最大吸力",
                "target_value": "40000Pa",
                "source_value": "22000Pa",
                "source_search_terms": ["最大吸力"],
                "target_search_terms": ["最大吸力"],
                "kb_intents": ["最大吸力"],
            }]
            candidates = server._clone_review_reuse_candidates(item, rules)
            by_id = {row["wiki_id"]: row for row in candidates}

            self.assertEqual(by_id["KB-NEW"]["priority"], "P1")
            self.assertEqual(by_id["KB-P2"]["priority"], "P2")
            self.assertNotIn("KB-CONFLICT", by_id)
            self.assertEqual(by_id["KB-P3-QUESTION"]["priority"], "P3")

            battery_rules = [{
                "feature_id": "battery_capacity",
                "feature_name": "电池容量",
                "target_value": "6400mAh",
                "source_value": "5200mAh",
                "source_search_terms": ["电池容量"],
                "target_search_terms": ["电池容量"],
                "kb_intents": ["电池容量"],
            }]
            battery_candidates = server._clone_review_reuse_candidates(item, battery_rules)
            battery_by_id = {row["wiki_id"]: row for row in battery_candidates}
            self.assertEqual(battery_by_id["KB-P3"]["priority"], "P3")

    def test_link_existing_adds_correct_relation_before_removing_clone(self):
        clone = self._clone("clone-op-link").get_json()
        detail = self.client.get(f"/api/matrix/clone-review/{clone['review_batch_id']}").get_json()
        item = detail["items"][0]
        decision = self.client.post(f"/api/matrix/clone-review/items/{item['item_id']}/decision", json={
            "decision": "link_existing",
            "selected_existing_wiki_id": "KB-NEW",
            "decision_reason": "同品类、同功能且目标参数一致",
        })
        self.assertEqual(decision.status_code, 200, decision.get_data(as_text=True))

        applied = self.client.post(f"/api/matrix/clone-review/{clone['review_batch_id']}/apply", json={})
        self.assertEqual(applied.status_code, 200, applied.get_data(as_text=True))
        completed = self.client.post(f"/api/matrix/clone-review/{clone['review_batch_id']}/complete", json={
            "manual_scope_confirmed": True,
            "rules_confirmed": True,
        })
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        self.assertEqual(completed.get_json()["status"], "ready_to_submit")
        expected = completed.get_json()["expected_submit_changes"]
        self.assertEqual({
            (row["question_wiki_id"], row["old_is_configured"], row["new_is_configured"])
            for row in expected
        }, {("KB-NEW", False, True)})

        with server.app.app_context():
            self.assertIsNone(server.ProductMatrix.query.filter_by(question_wiki_id="KB-NEW", product_name="P30 Pro").first())
            self.assertIsNone(server.ProductMatrix.query.filter_by(question_wiki_id="KB-OLD", product_name="P30 Pro").first())

        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            submitted = self.client.post("/api/matrix/submit_changes", json={
                "operation_id": "submit-link-final", "changes": expected,
            })
        self.assertEqual(submitted.status_code, 200, submitted.get_data(as_text=True))
        with server.app.app_context():
            correct = server.ProductMatrix.query.filter_by(question_wiki_id="KB-NEW", product_name="P30 Pro").one()
            self.assertTrue(correct.is_configured)
            self.assertIsNone(server.ProductMatrix.query.filter_by(question_wiki_id="KB-OLD", product_name="P30 Pro").first())

    def test_duplicate_source_kbs_linked_to_same_existing_kb_submit_once(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-OLD-DUPLICATE",
                product_name="P20",
                is_configured=True,
                question_content="最大吸力可以达到多少？",
                answer_content="最大吸力为 22000Pa。",
                product_category="扫地机",
            ))
            server.db.session.commit()

        clone = self._clone("clone-op-duplicate-source-link").get_json()
        batch_id = clone["review_batch_id"]
        items = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"]
        duplicate_items = [item for item in items if item["question_wiki_id"] in {"KB-OLD", "KB-OLD-DUPLICATE"}]
        self.assertEqual(len(duplicate_items), 2)
        for item in duplicate_items:
            decision = self.client.post(f"/api/matrix/clone-review/items/{item['item_id']}/decision", json={
                "decision": "link_existing",
                "selected_existing_wiki_id": "KB-NEW",
            })
            self.assertEqual(decision.status_code, 200, decision.get_data(as_text=True))

        applied = self.client.post(f"/api/matrix/clone-review/{batch_id}/apply", json={})
        self.assertEqual(applied.status_code, 200, applied.get_data(as_text=True))
        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={
            "manual_scope_confirmed": True,
            "rules_confirmed": True,
        })
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        expected = completed.get_json()["expected_submit_changes"]
        self.assertEqual(len(expected), 1)
        self.assertEqual(expected[0]["question_wiki_id"], "KB-NEW")

        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            submitted = self.client.post("/api/matrix/submit_changes", json={
                "operation_id": "submit-duplicate-source-link",
                "changes": expected,
            })
        self.assertEqual(submitted.status_code, 200, submitted.get_data(as_text=True))
        with server.app.app_context():
            self.assertEqual(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-NEW", product_name="P30 Pro", is_configured=True,
            ).count(), 1)
            self.assertIsNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P30 Pro",
            ).first())
            self.assertIsNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD-DUPLICATE", product_name="P30 Pro",
            ).first())

    def test_manual_id_can_be_verified_and_linked(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-MANUAL",
                product_name="P30-CANDIDATE-SOURCE",
                is_configured=True,
                question_content="集尘袋多久更换一次？",
                answer_content="建议根据使用频率定期更换。",
                product_category="扫地机",
            ))
            server.db.session.commit()
        clone = self._clone("clone-op-manual-id").get_json()
        item = self.client.get(f"/api/matrix/clone-review/{clone['review_batch_id']}").get_json()["items"][0]

        verified = self.client.post(
            f"/api/matrix/clone-review/items/{item['item_id']}/manual-candidate",
            json={"wiki_id": "KB-MANUAL"},
        )

        self.assertEqual(verified.status_code, 200, verified.get_data(as_text=True))
        candidate = next(row for row in verified.get_json()["item"]["reuse_candidates"] if row["wiki_id"] == "KB-MANUAL")
        self.assertTrue(candidate["manual_verified"])
        decision = self.client.post(f"/api/matrix/clone-review/items/{item['item_id']}/decision", json={
            "decision": "link_existing",
            "selected_existing_wiki_id": "KB-MANUAL",
        })
        self.assertEqual(decision.status_code, 200, decision.get_data(as_text=True))

    def test_manual_id_rejects_missing_disabled_and_cross_category_knowledge(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-OTHER-CATEGORY",
                product_name="OTHER-SOURCE",
                is_configured=True,
                question_content="如何安装？",
                answer_content="请按说明书安装。",
                product_category="洗地机",
            ))
            server.db.session.commit()
        clone = self._clone("clone-op-invalid-manual-id").get_json()
        item = self.client.get(f"/api/matrix/clone-review/{clone['review_batch_id']}").get_json()["items"][0]
        endpoint = f"/api/matrix/clone-review/items/{item['item_id']}/manual-candidate"

        missing = self.client.post(endpoint, json={"wiki_id": "KB-NOT-FOUND"})
        disabled = self.client.post(endpoint, json={"wiki_id": "KB-DISABLED"})
        cross_category = self.client.post(endpoint, json={"wiki_id": "KB-OTHER-CATEGORY"})

        self.assertEqual(missing.status_code, 422)
        self.assertEqual(disabled.status_code, 422)
        self.assertEqual(cross_category.status_code, 422)
        self.assertIn("品类不一致", cross_category.get_json()["message"])

    def test_link_existing_rejects_candidate_changed_after_scan(self):
        clone = self._clone("clone-op-stale-candidate").get_json()
        detail = self.client.get(f"/api/matrix/clone-review/{clone['review_batch_id']}").get_json()
        item = detail["items"][0]
        decision = self.client.post(f"/api/matrix/clone-review/items/{item['item_id']}/decision", json={
            "decision": "link_existing",
            "selected_existing_wiki_id": "KB-NEW",
        })
        self.assertEqual(decision.status_code, 200)
        with server.app.app_context():
            candidate = server.ProductMatrix.query.filter_by(question_wiki_id="KB-NEW").first()
            candidate.answer_content = "内容已变更为 50000Pa。"
            server.db.session.commit()

        applied = self.client.post(f"/api/matrix/clone-review/{clone['review_batch_id']}/apply", json={})

        self.assertEqual(applied.status_code, 207)
        self.assertIn("候选内容已变化", applied.get_json()["failed"][0]["message"])

    def test_high_medium_and_default_risk_flow(self):
        with server.app.app_context():
            server.db.session.add_all([
                server.ProductMatrix(
                    question_wiki_id="KB-MEDIUM",
                    product_name="P20",
                    is_configured=True,
                    question_content="最大吸力是否有提升？",
                    answer_content="该机型采用新一代风机。",
                    product_category="扫地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-DEFAULT",
                    product_name="P20",
                    is_configured=True,
                    question_content="尘盒如何清洁？",
                    answer_content="取出尘盒并清理灰尘。",
                    product_category="扫地机",
                ),
            ])
            server.db.session.commit()

        clone = self._clone("clone-op-risk-levels").get_json()
        detail = self.client.get(f"/api/matrix/clone-review/{clone['review_batch_id']}").get_json()
        items = {item["question_wiki_id"]: item for item in detail["items"]}

        self.assertEqual(items["KB-OLD"]["risk_level"], "high")
        self.assertTrue(items["KB-OLD"]["requires_attention"])
        self.assertEqual(items["KB-MEDIUM"]["risk_level"], "medium")
        self.assertTrue(items["KB-MEDIUM"]["requires_attention"])
        self.assertEqual(items["KB-DEFAULT"]["risk_level"], "default")
        self.assertEqual(items["KB-DEFAULT"]["decision"], "keep")
        self.assertEqual(items["KB-DEFAULT"]["apply_status"], "not_required")
        self.assertTrue(items["KB-DEFAULT"]["default_pass"])
        self.assertFalse(items["KB-DEFAULT"]["manually_decided"])
        self.assertFalse(items["KB-DEFAULT"]["requires_attention"])

        blocked = self.client.post(f"/api/matrix/clone-review/{clone['review_batch_id']}/complete", json={
            "manual_scope_confirmed": True,
        })
        self.assertEqual(blocked.status_code, 409)
        blocked_body = blocked.get_json()
        self.assertEqual(blocked_body["errors"], ["还有 2 条高、中风险或人工暂缓知识未形成最终决定"])
        self.assertNotIn("KB-OLD", "；".join(blocked_body["errors"]))

        for wiki_id in ("KB-OLD", "KB-MEDIUM"):
            decision = self.client.post(
                f"/api/matrix/clone-review/items/{items[wiki_id]['item_id']}/decision",
                json={"decision": "keep", "decision_reason": "人工确认可继续保留"},
            )
            self.assertEqual(decision.status_code, 200, decision.get_data(as_text=True))
            self.assertTrue(decision.get_json()["item"]["manually_decided"])

        completed = self.client.post(f"/api/matrix/clone-review/{clone['review_batch_id']}/complete", json={
            "manual_scope_confirmed": True,
            "rules_confirmed": True,
        })
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        self.assertEqual(completed.get_json()["status"], "ready_to_submit")

    def test_full_backup_exports_all_review_content_and_rules(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="export-backup-batch",
                operation_id="export-backup-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                scope_mode="all",
                scope_snapshot_json='{"exclude_tag_names":["不克隆"]}',
                strategy="append",
                status="reviewing",
                detection_methods_json='["manual_difference_ai"]',
                manual_difference_text="新品吸力变为 40000Pa",
                difference_rules_json='[{"feature_id":"suction"}]',
                source_count=2,
                item_count=2,
                created_by="admin",
            )
            server.db.session.add_all([
                batch,
                server.MatrixCloneReviewItem(
                    item_id="export-high-item",
                    batch_id=batch.batch_id,
                    question_wiki_id="KB-HIGH",
                    target_model="P30 Pro",
                    source_model="P20",
                    knowledge_revision="sha256:test-high",
                    question="最大吸力是多少？",
                    answer="=40000",
                    product_category="扫地机",
                    detection_status="scanned",
                    risk_level="high",
                    suggested_action="remove_target",
                    decision="pending",
                    decision_reason="等待人工确认",
                    difference_rule_ids_json='["export-rule"]',
                    evidence_sources_json='[{"source":"user_input","quote":"40000Pa"}]',
                    reuse_candidate_ids_json='[{"wiki_id":"KB-CANDIDATE"}]',
                    ai_impact_review_json='{"status":"completed","reason":"参数已变化"}',
                    apply_status="not_applied",
                ),
                server.MatrixCloneReviewItem(
                    item_id="export-default-item",
                    batch_id=batch.batch_id,
                    question_wiki_id="KB-DEFAULT",
                    target_model="P30 Pro",
                    source_model="P20",
                    question="尘盒如何清洁？",
                    answer="取出尘盒并清理灰尘。",
                    product_category="扫地机",
                    detection_status="scanned",
                    risk_level="default",
                    suggested_action="keep",
                    decision="keep",
                    apply_status="not_required",
                ),
                server.MatrixCloneDifferenceRule(
                    rule_id="export-rule",
                    batch_id=batch.batch_id,
                    target_model="P30 Pro",
                    feature_id="suction",
                    feature_name="最大吸力",
                    source_value="22000Pa",
                    target_value="40000Pa",
                    difference_type="value_changed",
                    evidence_sources_json='["user_input"]',
                    evidence_quote="新品吸力变为 40000Pa",
                    status="confirmed",
                ),
            ])
            server.db.session.commit()

        response = self.client.get("/api/matrix/clone-review/export-backup-batch/export.xlsx")

        self.assertEqual(response.status_code, 200)
        self.assertIn("clone_review_backup_export-b", response.headers["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertEqual(workbook.sheetnames, ["批次信息", "需人工处理知识", "差异规则"])

        batch_values = {
            row[0].value: row[1].value
            for row in workbook["批次信息"].iter_rows(min_row=2)
        }
        self.assertEqual(batch_values["batch_id"], "export-backup-batch")
        self.assertEqual(batch_values["scope_snapshot_json"], '{"exclude_tag_names":["不克隆"]}')

        item_sheet = workbook["需人工处理知识"]
        item_headers = [cell.value for cell in item_sheet[1]]
        item_rows = {
            row[item_headers.index("question_wiki_id")].value: row
            for row in item_sheet.iter_rows(min_row=2)
        }
        self.assertEqual(set(item_rows), {"KB-HIGH", "KB-DEFAULT"})
        high_row = item_rows["KB-HIGH"]
        answer_cell = high_row[item_headers.index("answer")]
        self.assertEqual(answer_cell.value, "=40000")
        self.assertEqual(answer_cell.data_type, "s")
        self.assertIn("40000Pa", high_row[item_headers.index("evidence_sources_json")].value)

        rule_sheet = workbook["差异规则"]
        rule_headers = [cell.value for cell in rule_sheet[1]]
        rule_row = next(rule_sheet.iter_rows(min_row=2))
        self.assertEqual(rule_row[rule_headers.index("rule_id")].value, "export-rule")
        self.assertEqual(rule_row[rule_headers.index("target_value")].value, "40000Pa")

    def test_ai_impact_review_only_auto_passes_unrelated_context_with_source_evidence(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="impact-review-batch",
                operation_id="impact-review-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json='["manual_difference_ai"]',
            )
            unrelated = server.MatrixCloneReviewItem(
                item_id="impact-unrelated",
                batch_id=batch.batch_id,
                question_wiki_id="KB-SURFACE",
                source_model="P20",
                target_model="P30 Pro",
                question="为什么机器人机身会有不规则的痕迹？",
                answer="平整的痕迹为注塑工艺产生，不影响正常使用。",
                product_category="扫地机",
            )
            affected = server.MatrixCloneReviewItem(
                item_id="impact-affected",
                batch_id=batch.batch_id,
                question_wiki_id="KB-HEIGHT",
                source_model="P20",
                target_model="P30 Pro",
                question="机身高度会影响能否进入床底吗？",
                answer="会影响低矮空间的通行能力。",
                product_category="扫地机",
            )
            invalid_evidence = server.MatrixCloneReviewItem(
                item_id="impact-invalid-evidence",
                batch_id=batch.batch_id,
                question_wiki_id="KB-INVALID",
                source_model="P20",
                target_model="P30 Pro",
                question="机身外壳如何清洁？",
                answer="使用软布擦拭外壳。",
                product_category="扫地机",
            )
            direct_high = server.MatrixCloneReviewItem(
                item_id="impact-direct-high",
                batch_id=batch.batch_id,
                question_wiki_id="KB-DIRECT-HIGH",
                source_model="P20",
                target_model="P30 Pro",
                question="机身高度是多少？",
                answer="机身高度为 11.9cm。",
                product_category="扫地机",
            )
            rule = server.MatrixCloneDifferenceRule(
                rule_id="impact-height-rule",
                batch_id=batch.batch_id,
                target_model="P30 Pro",
                feature_id="body_height",
                feature_name="机身高度",
                source_value="11.9cm",
                target_value="8.98cm",
                difference_type="value_changed",
                source_search_terms_json='["机身高度", "机身"]',
                kb_intents_json='["机身高度"]',
                evidence_sources_json='["user_input"]',
            )
            server.db.session.add_all([batch, unrelated, affected, invalid_evidence, direct_high, rule])
            server.db.session.commit()

            def impact_response(_config, _system_prompt, user_prompt, temperature=0.0):
                payload = __import__("json").loads(user_prompt.split("：", 1)[1])
                self.assertEqual({row["item_id"] for row in payload}, {
                    "impact-unrelated", "impact-affected", "impact-invalid-evidence", "impact-direct-high",
                })
                return __import__("json").dumps({"reviews": [
                    {
                        "item_id": "impact-unrelated",
                        "relation": "unrelated",
                        "evidence": "不规则的痕迹",
                        "reason": "知识讨论外观痕迹和注塑工艺，不讨论高度。",
                    },
                    {
                        "item_id": "impact-affected",
                        "relation": "affected",
                        "evidence": "机身高度会影响能否进入床底吗？",
                        "reason": "知识结论直接讨论低矮空间通行能力。",
                    },
                    {
                        "item_id": "impact-invalid-evidence",
                        "relation": "unrelated",
                        "evidence": "不存在的原文证据",
                        "reason": "无效引用不得自动放行。",
                    },
                    {
                        "item_id": "impact-direct-high",
                        "relation": "unrelated",
                        "evidence": "机身高度为 11.9cm",
                        "reason": "即使 AI 判断不相关，直接数值冲突也不能自动降级。",
                    },
                ]}, ensure_ascii=False)

            with patch.object(server, "load_ai_config", return_value={
                "api_key": "configured", "base_url": "http://ai.test", "model": "test-model",
            }), patch.object(server, "_ai_call_llm", side_effect=impact_response):
                server._scan_clone_review_batch(batch)
                server.db.session.commit()

            rows = {item.item_id: item for item in server.MatrixCloneReviewItem.query.all()}
            self.assertEqual(rows["impact-unrelated"].risk_level, "default")
            self.assertEqual(rows["impact-unrelated"].decision, "keep")
            self.assertEqual(
                server._clone_review_json(rows["impact-unrelated"].ai_impact_review_json, {})["relation"],
                "unrelated",
            )
            self.assertEqual(rows["impact-affected"].risk_level, "medium")
            self.assertEqual(rows["impact-direct-high"].risk_level, "high")
            self.assertEqual(
                server._clone_review_json(rows["impact-direct-high"].ai_impact_review_json, {})["risk_effect"],
                "kept_high_due_to_direct_evidence",
            )
            self.assertEqual(rows["impact-invalid-evidence"].risk_level, "medium")
            self.assertEqual(
                server._clone_review_json(rows["impact-invalid-evidence"].ai_impact_review_json, {})["relation"],
                "uncertain",
            )

    def test_ai_impact_review_uses_micro_batches_of_at_most_five_items(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="impact-micro-batch",
                operation_id="impact-micro-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json='["manual_difference_ai"]',
            )
            rule = server.MatrixCloneDifferenceRule(
                rule_id="impact-micro-rule",
                batch_id=batch.batch_id,
                target_model="P30 Pro",
                feature_id="body_height",
                feature_name="机身高度",
                source_value="11.9cm",
                target_value="8.98cm",
                difference_type="value_changed",
                evidence_sources_json='["user_input"]',
            )
            items = [
                server.MatrixCloneReviewItem(
                    item_id=f"impact-micro-{index}",
                    batch_id=batch.batch_id,
                    question_wiki_id=f"KB-MICRO-{index}",
                    source_model="P20",
                    target_model="P30 Pro",
                    question=f"第 {index} 条机身高度知识",
                    answer="机身高度为 11.9cm。",
                    product_category="扫地机",
                )
                for index in range(6)
            ]
            server.db.session.add_all([batch, rule, *items])
            server.db.session.commit()

            batch_sizes = []

            def review_micro_batch(review_items, _rules):
                batch_sizes.append(len(review_items))
                return {}, "模拟 AI 不可用"

            with patch.object(server, "_clone_review_ai_impact_reviews", side_effect=review_micro_batch):
                server._scan_clone_review_batch(batch, commit_every=1)

            self.assertEqual(batch_sizes, [5, 1])
            self.assertTrue(all(size <= 5 for size in batch_sizes))

    def test_legacy_unmatched_item_is_normalized_to_default(self):
        item = server.MatrixCloneReviewItem(
            item_id="legacy-item",
            batch_id="legacy-batch",
            question_wiki_id="KB-LEGACY",
            target_model="P30 Pro",
            source_model="P20",
            detection_status="scanned",
            risk_level="unknown",
            suggested_action="defer",
            decision="pending",
            apply_status="not_applied",
            difference_rule_ids_json="[]",
        )

        changed = server._normalize_clone_review_default_item(item)

        self.assertTrue(changed)
        self.assertEqual(item.risk_level, "default")
        self.assertEqual(item.decision, "keep")
        self.assertEqual(item.apply_status, "not_required")
        self.assertEqual(item.result_wiki_id, "KB-LEGACY")

    def test_zero_pending_manual_batch_waits_for_scope_confirmation(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-OLD",
                product_name="P30 Pro",
                is_configured=True,
                question_content="最大吸力是多少？",
                answer_content="最大吸力为 22000Pa。",
                product_category="扫地机",
            ))
            server.db.session.commit()
        clone = self._clone("clone-op-zero-pending").get_json()
        self.assertEqual(clone["pending_count"], 0)
        batch_id = clone["review_batch_id"]
        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
        self.assertEqual(detail["status"], "reviewing")
        self.assertEqual(detail["rules"][0]["status"], "needs_confirmation")

        missing_confirmation = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={})
        self.assertEqual(missing_confirmation.status_code, 422)

        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={
            "manual_scope_confirmed": True,
            "rules_confirmed": True,
        })
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        self.assertEqual(completed.get_json()["status"], "ready_to_submit")
        self.assertEqual(completed.get_json()["rules"][0]["status"], "valid")

    def test_dual_mode_requires_explicit_confirmation_for_unmatched_manual_rule(self):
        parameter_rules = [{
            "question_wiki_id": "KB-OLD", "target_model": "P30 Pro", "feature_id": "max_suction",
            "feature_name": "最大吸力", "source_value": "22000Pa", "target_value": "40000Pa",
            "difference_type": "value_changed", "evidence_sources": ["parameter_check"], "status": "valid",
        }]
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "_clone_review_parameter_difference_rules", return_value=(parameter_rules, None)
        ), patch.object(server, "load_ai_config", return_value={}), patch.object(
            server, "_refresh_clone_parameter_comparison", return_value={
                "snapshot_id": "snapshot-test-latest", "source_version": "v-test-latest",
                "comparison_run_id": "run-test-latest", "scan_run_id": "scan-test-latest",
                "result_counts": {},
            }
        ):
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-dual-confirm",
                "mode": "model", "source": "P20", "targets": ["P30 Pro"], "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True, "auto_scan": True,
                    "detection_methods": ["parameter_check", "manual_difference_ai"],
                    "manual_difference_text": "旧型号最大吸力 22000Pa，新品 40000Pa；旧型号喷淋 8 孔，新品 16 孔",
                },
            })
        batch_id = response.get_json()["review_batch_id"]
        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
        item = detail["items"][0]
        self.client.post(f"/api/matrix/clone-review/items/{item['item_id']}/decision", json={"decision": "keep"})

        blocked = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={})
        self.assertEqual(blocked.status_code, 422)
        self.assertEqual(blocked.get_json()["code"], "difference_rule_confirmation_required")

        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={"rules_confirmed": True})
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        self.assertEqual(completed.get_json()["status"], "ready_to_submit")

    def test_dual_mode_returns_parameter_difference_rules_with_feature_context(self):
        parameter_rules = [{
            "source_model": "P20", "target_model": "P30 Pro", "feature_id": "max_suction",
            "feature_name": "最大吸力", "feature_description": "用于说明吸入灰尘的峰值能力",
            "source_value": "22000Pa", "target_value": "40000Pa",
            "difference_type": "value_changed", "evidence_sources": ["parameter_check"], "status": "valid",
        }]
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "_clone_review_parameter_difference_rules", return_value=(parameter_rules, None)
        ), patch.object(server, "load_ai_config", return_value={}), patch.object(
            server, "_refresh_clone_parameter_comparison", return_value={
                "snapshot_id": "snapshot-test-latest", "source_version": "v-test-latest",
                "comparison_run_id": "run-test-latest", "scan_run_id": "scan-test-latest",
                "result_counts": {},
            }
        ):
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-dual-rule-display",
                "mode": "model", "source": "P20", "targets": ["P30 Pro"], "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True, "auto_scan": True,
                    "detection_methods": ["parameter_check", "manual_difference_ai"],
                    "manual_difference_text": "旧型号最大吸力 22000Pa，新品 40000Pa",
                },
            })
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            detail = self.client.get(f"/api/matrix/clone-review/{response.get_json()['review_batch_id']}").get_json()

        self.assertEqual(len(detail["rules"]), 1)
        self.assertEqual(len(detail["parameter_rules"]), 1)
        self.assertEqual(detail["parameter_rules"][0]["difference_type"], "value_changed")
        self.assertEqual(detail["parameter_rules"][0]["feature_description"], "用于说明吸入灰尘的峰值能力")
        self.assertIn(detail["rules"][0]["rule_id"], detail["items"][0]["difference_rule_ids"])
        self.assertIn(detail["parameter_rules"][0]["rule_id"], detail["items"][0]["difference_rule_ids"])

    def test_link_existing_batch_is_submitted_only_after_all_expected_changes(self):
        clone = self._clone("clone-op-submit-gate").get_json()
        batch_id = clone["review_batch_id"]
        item = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]
        self.client.post(f"/api/matrix/clone-review/items/{item['item_id']}/decision", json={
            "decision": "link_existing", "selected_existing_wiki_id": "KB-NEW",
        })
        self.client.post(f"/api/matrix/clone-review/{batch_id}/apply", json={})
        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={
            "manual_scope_confirmed": True, "rules_confirmed": True,
        }).get_json()
        expected = completed["expected_submit_changes"]

        self.assertEqual(len(expected), 1)
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            submitted = self.client.post("/api/matrix/submit_changes", json={
                "operation_id": "submit-link-final", "changes": expected,
            })
        self.assertEqual(submitted.status_code, 200, submitted.get_data(as_text=True))
        self.assertEqual(self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["status"], "submitted")

    def test_remove_target_draft_never_writes_target_relation(self):
        clone = self._clone("clone-op-remove-draft").get_json()
        batch_id = clone["review_batch_id"]
        item = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]
        decision = self.client.post(f"/api/matrix/clone-review/items/{item['item_id']}/decision", json={
            "decision": "remove_target", "decision_reason": "确认新品不应关联此知识",
        })
        self.assertEqual(decision.status_code, 200)
        applied = self.client.post(f"/api/matrix/clone-review/{batch_id}/apply", json={})
        self.assertEqual(applied.status_code, 200, applied.get_data(as_text=True))
        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={
            "manual_scope_confirmed": True, "rules_confirmed": True,
        })
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        self.assertEqual(completed.get_json()["expected_submit_changes"], [])
        with server.app.app_context():
            self.assertIsNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P30 Pro",
            ).first())

    def test_review_batch_list_returns_recent_statuses(self):
        batch_id = self._clone("clone-op-list").get_json()["review_batch_id"]
        response = self.client.get("/api/matrix/clone-review-batches?status=reviewing,ready_to_submit,submitted")
        self.assertEqual(response.status_code, 200)
        rows = response.get_json()["batches"]
        self.assertEqual(rows[0]["batch_id"], batch_id)
        self.assertEqual(rows[0]["status"], "reviewing")
        self.assertTrue(rows[0]["can_delete"])

    def test_deletes_unsubmitted_review_batch_and_children(self):
        batch_id = self._clone("clone-op-delete-unused").get_json()["review_batch_id"]
        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
        self.assertTrue(detail["can_delete"])

        response = self.client.delete(f"/api/matrix/clone-review/{batch_id}")

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(len(response.get_json()["removed_changes"]), 1)
        with server.app.app_context():
            self.assertIsNone(server.db.session.get(server.MatrixCloneReviewBatch, batch_id))
            self.assertEqual(server.MatrixCloneDifferenceRule.query.filter_by(batch_id=batch_id).count(), 0)
            self.assertEqual(server.MatrixCloneReviewItem.query.filter_by(batch_id=batch_id).count(), 0)

    def test_deletes_submitted_review_batch_without_reverting_matrix_data(self):
        batch_id = self._clone("clone-op-delete-submitted").get_json()["review_batch_id"]
        with server.app.app_context():
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            batch.status = "submitted"
            batch.submitted_at = server.datetime.utcnow()
            server.db.session.commit()
            target_count = server.ProductMatrix.query.filter_by(product_name="P30 Pro").count()

        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
        self.assertTrue(detail["can_delete"])

        response = self.client.delete(f"/api/matrix/clone-review/{batch_id}")

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertTrue(response.get_json()["submitted_data_preserved"])
        with server.app.app_context():
            self.assertIsNone(server.db.session.get(server.MatrixCloneReviewBatch, batch_id))
            self.assertEqual(
                server.ProductMatrix.query.filter_by(product_name="P30 Pro").count(),
                target_count,
            )

    def test_rejects_deleting_running_review_batch(self):
        batch_id = self._clone("clone-op-delete-running").get_json()["review_batch_id"]
        with server.app.app_context():
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            batch.status = "scanning"
            server.db.session.commit()
        self.assertTrue(server._claim_matrix_clone_scan(batch_id))
        try:
            detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
            self.assertFalse(detail["can_delete"])
            response = self.client.delete(f"/api/matrix/clone-review/{batch_id}")
        finally:
            server._release_matrix_clone_scan(batch_id)

        self.assertEqual(response.status_code, 409)
        with server.app.app_context():
            self.assertIsNotNone(server.db.session.get(server.MatrixCloneReviewBatch, batch_id))

    def test_deletes_interrupted_review_batch_that_can_continue(self):
        batch_id = self._clone("clone-op-delete-interrupted").get_json()["review_batch_id"]
        with server.app.app_context():
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            batch.status = "scanning"
            server.db.session.commit()

        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
        self.assertFalse(detail["scan_active"])
        self.assertTrue(detail["can_delete"])

        response = self.client.delete(f"/api/matrix/clone-review/{batch_id}")

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        with server.app.app_context():
            self.assertIsNone(server.db.session.get(server.MatrixCloneReviewBatch, batch_id))

    def test_rejects_deleting_another_users_review_batch(self):
        batch_id = self._clone("clone-op-delete-other-user").get_json()["review_batch_id"]
        with server.app.app_context():
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            batch.created_by = "another-user"
            server.db.session.commit()

        response = self.client.delete(f"/api/matrix/clone-review/{batch_id}")

        self.assertEqual(response.status_code, 403)
        with server.app.app_context():
            self.assertIsNotNone(server.db.session.get(server.MatrixCloneReviewBatch, batch_id))

    def test_dual_source_conflict_is_blocked(self):
        parameter_rules = [{
            "question_wiki_id": "KB-OLD",
            "target_model": "P30 Pro",
            "feature_id": "max_suction",
            "feature_name": "最大吸力",
            "source_value": "22000Pa",
            "target_value": "50000Pa",
            "difference_type": "value_changed",
            "source_search_terms": ["22000Pa", "最大吸力"],
            "target_search_terms": ["50000Pa", "最大吸力"],
            "kb_intents": ["最大吸力是多少"],
            "evidence_sources": ["parameter_check"],
            "status": "valid",
        }]
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "_clone_review_parameter_difference_rules", return_value=(parameter_rules, None)
        ), patch.object(server, "load_ai_config", return_value={}), patch.object(
            server, "_refresh_clone_parameter_comparison", return_value={
                "snapshot_id": "snapshot-test-latest", "source_version": "v-test-latest",
                "comparison_run_id": "run-test-latest", "scan_run_id": "scan-test-latest",
                "result_counts": {},
            }
        ):
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-source-conflict",
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True,
                    "auto_scan": True,
                    "detection_methods": ["parameter_check", "manual_difference_ai"],
                    "manual_difference_text": "旧型号最大吸力 22000Pa，新品 40000Pa",
                },
            })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        batch_id = response.get_json()["review_batch_id"]
        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
        self.assertEqual(detail["items"][0]["risk_level"], "default")
        self.assertEqual(detail["items"][0]["detection_status"], "scan_failed")
        self.assertEqual(detail["items"][0]["suggested_action"], "defer")
        self.assertIn("来源冲突", detail["error_message"])

        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={})
        self.assertEqual(completed.status_code, 409)
        self.assertEqual(completed.get_json()["code"], "clone_review_incomplete")

    def test_parameter_catalog_audit_marks_consistent_suspected_and_missing_values(self):
        manual_rules = [
            {"feature_id": "body_height", "feature_name": "机身高度", "source_value": "11.9cm", "target_value": "8.98cm"},
            {"feature_id": "max_suction", "feature_name": "最大吸力", "source_value": "22000Pa", "target_value": "40000Pa"},
            {"feature_id": "suction_power", "feature_name": "吸力", "source_value": "22000Pa", "target_value": "40000Pa"},
            {"feature_id": "spray_hole_count", "feature_name": "喷淋孔数", "source_value": "8孔", "target_value": "16孔"},
        ]
        snapshots = {
            "P20 Ultra活水版": ({"机身高度": "机身高度为119mm", "最大吸力": "22000Pa"}, {}, None),
            "P30 Pro": ({"机身高度": "机身高度为89.8mm", "最大吸力": "36000pa"}, {}, None),
        }
        with patch.object(server, "_clone_review_catalog_snapshot", side_effect=lambda model: snapshots[model]):
            rows, error = server._clone_review_parameter_catalog_audit(
                manual_rules, "P20 Ultra活水版", "P30 Pro",
            )

        self.assertIsNone(error)
        by_feature = {row["feature_id"]: row for row in rows}
        self.assertEqual(by_feature["body_height"]["result"], "consistent")
        self.assertEqual(by_feature["max_suction"]["source_result"], "consistent")
        self.assertEqual(by_feature["max_suction"]["target_result"], "catalog_suspected")
        self.assertEqual(by_feature["suction_power"]["catalog_feature_name"], "最大吸力")
        self.assertEqual(by_feature["suction_power"]["target_result"], "catalog_suspected")
        self.assertEqual(by_feature["spray_hole_count"]["result"], "missing_parameter")

    def test_parameter_difference_rules_ignore_placeholders_and_keep_plain_description(self):
        source_values = {"最大吸力": "22000Pa", "清洁液盒容量": "600ml"}
        target_values = {"最大吸力": "40000Pa", "清洁液盒容量": "暂无数据"}
        features = {
            "最大吸力": {"feature_id": "max_suction", "description": "最大吸力说明"},
            "清洁液盒容量": {"feature_id": "solution", "description": "容量说明"},
        }
        with patch.object(
            server,
            "_clone_review_catalog_model_data",
            side_effect=[(source_values, features, None), (target_values, features, None)],
        ):
            rules, error = server._clone_review_parameter_difference_rules("P20", "P30 Pro")

        self.assertIsNone(error)
        self.assertEqual(len(rules), 1)
        self.assertEqual(rules[0]["feature_name"], "最大吸力")
        self.assertEqual(rules[0]["feature_description"], "最大吸力说明")

    def test_parameter_difference_rules_keep_boolean_and_extended_value_changes(self):
        source_values = {
            "脏污检测": "支持",
            "一次性滤网": "有",
            "工作模式切换": "自动/强力/轻拖/吸水",
        }
        target_values = {
            "脏污检测": "不支持",
            "一次性滤网": "没有",
            "工作模式切换": "自动/强力/轻拖/吸水，自动模式可在 APP 中调节",
        }
        features = {
            name: {"feature_id": name, "description": f"{name}说明"}
            for name in source_values
        }
        with patch.object(
            server,
            "_clone_review_catalog_model_data",
            side_effect=[(source_values, features, None), (target_values, features, None)],
        ):
            rules, error = server._clone_review_parameter_difference_rules("旧型号", "新型号")

        self.assertIsNone(error)
        self.assertEqual({rule["feature_name"] for rule in rules}, set(source_values))

    def test_catalog_value_matching_does_not_hide_boolean_negation(self):
        self.assertFalse(server._clone_review_catalog_value_matches("脏污检测", "支持", "不支持"))
        self.assertFalse(server._clone_review_catalog_value_matches("一次性滤网", "有", "没有"))
        self.assertTrue(server._clone_review_catalog_value_matches("一次性滤网", "支持", "支持一次性滤网"))

    def test_boolean_values_are_not_standalone_knowledge_evidence(self):
        self.assertFalse(server._clone_review_value_is_direct_evidence("支持"))
        self.assertFalse(server._clone_review_value_is_direct_evidence("不支持"))
        self.assertFalse(server._clone_review_value_is_direct_evidence("有"))
        self.assertFalse(server._clone_review_value_is_direct_evidence("没有"))
        self.assertTrue(server._clone_review_value_is_direct_evidence("支持一次性滤网"))
        self.assertTrue(server._clone_review_value_is_direct_evidence("25000Pa"))

    def test_dirt_detection_terms_cover_reversed_and_indicator_wording(self):
        terms = server.difference_feature_terms(
            "feature-uuid",
            "脏污检测",
            "实时感知地面及滚刷脏污程度，动态调节清洁水量。",
        )
        self.assertIn("脏污检测", terms)
        self.assertIn("检测脏污", terms)
        self.assertIn("污渍检测", terms)
        self.assertIn("脏污指示", terms)
        self.assertIn("脏污程度", terms)

    def test_feature_terms_use_shared_semantics_for_unregistered_feature(self):
        terms = server.difference_feature_terms(
            "generic-obstacle-feature",
            "障碍物检测",
            "实时检测前方障碍物状态并调整运行方式。",
        )

        self.assertIn("检测障碍物", terms)
        self.assertIn("障碍物识别", terms)
        self.assertIn("障碍物感知", terms)
        self.assertIn("障碍物状态", terms)

    def test_catalog_description_derived_term_participates_in_context_scan(self):
        rule = {
            "feature_id": "generic-obstacle-feature",
            "feature_name": "障碍物检测",
            "feature_description": "实时检测前方障碍物状态并调整运行方式。",
            "source_value": "支持",
            "target_value": "不支持",
            "source_search_terms": [],
            "kb_intents": [],
        }

        self.assertTrue(server._clone_review_rule_matches_context("设备会展示障碍物状态。", rule))

    def test_boolean_parameter_rule_requires_feature_context_during_scan(self):
        with server.app.app_context():
            server.db.session.add_all([
                server.ProductMatrix(
                    question_wiki_id="KB-DIRT-DETECTION",
                    product_name="P20",
                    is_configured=True,
                    question_content="每个档位都可以智能检测脏污度？",
                    answer_content="自动档会根据脏污程度调节工作状态。",
                    product_category="洗地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-GENERIC-SUPPORT",
                    product_name="P20",
                    is_configured=True,
                    question_content="清洁液如何使用？",
                    answer_content="本机支持按说明添加清洁液。",
                    product_category="洗地机",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-GENERIC-NOT-SUPPORT",
                    product_name="P20",
                    is_configured=True,
                    question_content="是否可以查看地图？",
                    answer_content="本机暂不支持地图功能。",
                    product_category="洗地机",
                ),
            ])
            server.db.session.commit()

        parameter_rules = [{
            "source_model": "P20", "target_model": "P30 Pro", "feature_id": "dirt_detection",
            "feature_name": "脏污检测", "source_value": "支持", "target_value": "不支持",
            "difference_type": "value_changed", "source_search_terms": ["脏污检测", "支持", "不支持"],
            "target_search_terms": ["脏污检测", "不支持"], "kb_intents": ["脏污检测"],
            "evidence_sources": ["parameter_check"], "status": "valid",
        }]

        def affected_reviews(items, _rules):
            return ({item.item_id: {
                "status": "completed", "relation": "affected",
                "evidence": item.question, "reason": "知识直接描述该功能。",
            } for item in items}, None)

        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "_clone_review_parameter_difference_rules", return_value=(parameter_rules, None)
        ), patch.object(server, "_clone_review_ai_impact_reviews", side_effect=affected_reviews), patch.object(
            server, "_refresh_clone_parameter_comparison", return_value={
                "snapshot_id": "snapshot-test-latest", "source_version": "v-test-latest",
                "comparison_run_id": "run-test-latest", "scan_run_id": "scan-test-latest",
                "result_counts": {},
            }
        ):
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-boolean-feature-context",
                "mode": "model", "source": "P20", "targets": ["P30 Pro"], "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True, "auto_scan": True,
                    "detection_methods": ["parameter_check"],
                },
            })
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            detail = self.client.get(
                f"/api/matrix/clone-review/{response.get_json()['review_batch_id']}"
            ).get_json()

        items = {item["question_wiki_id"]: item for item in detail["items"]}
        self.assertEqual(items["KB-DIRT-DETECTION"]["risk_level"], "medium")
        self.assertEqual(len(items["KB-DIRT-DETECTION"]["difference_rule_ids"]), 1)
        self.assertEqual(items["KB-GENERIC-SUPPORT"]["risk_level"], "default")
        self.assertEqual(items["KB-GENERIC-SUPPORT"]["difference_rule_ids"], [])
        self.assertEqual(items["KB-GENERIC-NOT-SUPPORT"]["risk_level"], "default")
        self.assertEqual(items["KB-GENERIC-NOT-SUPPORT"]["difference_rule_ids"], [])

    def test_source_conflict_comparison_normalizes_equivalent_units(self):
        manual = {
            "feature_id": "body_height", "feature_name": "机身高度",
            "source_value": "11.9cm", "target_value": "8.98cm", "evidence_sources": ["user_input"],
        }
        parameter = {
            "feature_id": "body_height", "feature_name": "机身高度",
            "source_value": "11.9cm", "target_value": "89.8mm", "evidence_sources": ["parameter_check"],
        }
        self.assertFalse(server._clone_review_has_source_conflict([manual, parameter]))

    def test_object_recognition_comparison_ignores_other_numbers(self):
        self.assertTrue(server._clone_review_catalog_value_matches(
            "避障方式",
            "超 300 + 种物体识别",
            "前结构光 + 沿边结构光 2.0 超 300 种视觉识别",
        ))

    def test_review_validation_has_stable_error_codes(self):
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            missing_method = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "missing-method",
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "review": {"enabled": True, "detection_methods": []},
            })
            missing_text = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "missing-text",
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "review": {"enabled": True, "detection_methods": ["manual_difference_ai"]},
            })
        self.assertEqual(missing_method.status_code, 422)
        self.assertEqual(missing_method.get_json()["code"], "detection_method_required")
        self.assertEqual(missing_text.status_code, 422)
        self.assertEqual(missing_text.get_json()["code"], "manual_difference_required")

    def test_clone_rejects_invalid_inputs_before_creating_a_target_column(self):
        invalid_requests = [
            ({"mode": "unknown", "source": "P20", "targets": ["P30 Pro"]}, 400),
            ({"mode": "model", "source": "P20", "targets": "P30 Pro"}, 400),
            ({"mode": "model", "source": "P20", "targets": ["P30 Pro", "P30 Pro"]}, 400),
            ({"mode": "model", "source": "P20", "targets": ["P20"]}, 400),
            ({"mode": "model", "source": "P20", "targets": ["NOT-IN-CATALOG"]}, 422),
        ]
        for payload, expected_status in invalid_requests:
            with self.subTest(payload=payload), patch.object(
                server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None
            ):
                response = self.client.post("/api/matrix/clone_config", json=payload)
            self.assertEqual(response.status_code, expected_status, response.get_data(as_text=True))

        with server.app.app_context():
            self.assertIsNone(server.MatrixColumn.query.filter_by(product_name="NOT-IN-CATALOG").first())

    def test_review_clone_with_no_available_source_data_does_not_create_a_batch(self):
        with server.app.app_context():
            source = server.ProductMatrix.query.filter_by(question_wiki_id="KB-OLD", product_name="P20").one()
            source.is_configured = False
            server.db.session.commit()

        response = self._clone("clone-op-empty-source")

        self.assertEqual(response.status_code, 422, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["code"], "clone_review_source_empty")
        with server.app.app_context():
            self.assertEqual(server.MatrixCloneReviewBatch.query.count(), 0)

    def test_category_fallback_ignores_unconfigured_rows(self):
        with server.app.app_context():
            server.db.session.add_all([
                server.ProductMatrix(
                    question_wiki_id="KB-CATEGORY-CONFIGURED",
                    product_name="P20",
                    is_configured=True,
                    product_category="回收测试",
                ),
                server.ProductMatrix(
                    question_wiki_id="KB-CATEGORY-DISABLED",
                    product_name="P20",
                    is_configured=False,
                    product_category="回收测试",
                ),
            ])
            server.db.session.commit()

            source_items, _ = server._resolve_matrix_clone_source_items("category", "回收测试")

        self.assertEqual(set(source_items), {"KB-CATEGORY-CONFIGURED"})

    def test_complete_marks_source_changes_stale_and_returns_to_reviewing(self):
        clone = self._clone("clone-op-stale-complete").get_json()
        batch_id = clone["review_batch_id"]
        item = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]
        self.assertEqual(self.client.post(
            f"/api/matrix/clone-review/items/{item['item_id']}/decision",
            json={"decision": "keep"},
        ).status_code, 200)
        with server.app.app_context():
            source = server.ProductMatrix.query.filter_by(question_wiki_id="KB-OLD", product_name="P20").one()
            source.answer_content = "最大吸力内容已更新。"
            server.db.session.commit()

        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={
            "manual_scope_confirmed": True, "rules_confirmed": True,
        })

        self.assertEqual(completed.status_code, 409, completed.get_data(as_text=True))
        self.assertEqual(completed.get_json()["code"], "knowledge_revision_stale")
        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()
        self.assertEqual(detail["status"], "reviewing")
        self.assertEqual(detail["items"][0]["detection_status"], "stale")
        self.assertEqual(detail["items"][0]["apply_status"], "stale")

    def test_refresh_stale_item_uses_latest_source_and_preserves_other_decisions(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-SECOND",
                product_name="P20",
                is_configured=True,
                question_content="如何清洁尘盒？",
                answer_content="取出尘盒后清洁。",
                product_category="扫地机",
            ))
            server.db.session.commit()
        clone = self._clone("clone-op-refresh-stale-item").get_json()
        batch_id = clone["review_batch_id"]
        items = {
            item["question_wiki_id"]: item
            for item in self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"]
        }
        second = self.client.post(
            f"/api/matrix/clone-review/items/{items['KB-SECOND']['item_id']}/decision",
            json={"decision": "defer", "decision_reason": "保留另一条人工判断"},
        )
        self.assertEqual(second.status_code, 200, second.get_data(as_text=True))
        with server.app.app_context():
            source = server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P20",
            ).one()
            source.question_content = "当前知识标题已同步更新？"
            source.answer_content = "当前知识答案已同步更新。"
            server.db.session.commit()
        stale = self.client.post(
            f"/api/matrix/clone-review/items/{items['KB-OLD']['item_id']}/decision",
            json={"decision": "defer", "decision_reason": "旧判断"},
        )
        self.assertEqual(stale.status_code, 409, stale.get_data(as_text=True))

        with patch.object(server, "load_ai_config", return_value={}):
            refreshed = self.client.post(
                f"/api/matrix/clone-review/items/{items['KB-OLD']['item_id']}/refresh",
            )

        self.assertEqual(refreshed.status_code, 200, refreshed.get_data(as_text=True))
        body = refreshed.get_json()
        self.assertTrue(body["success"])
        self.assertEqual(body["item"]["question"], "当前知识标题已同步更新？")
        self.assertEqual(body["item"]["answer"], "当前知识答案已同步更新。")
        self.assertEqual(body["item"]["detection_status"], "scanned")
        self.assertEqual(body["item"]["decision_reason"], "")
        refreshed_items = {item["question_wiki_id"]: item for item in body["batch"]["items"]}
        self.assertEqual(refreshed_items["KB-SECOND"]["decision"], "defer")
        self.assertEqual(refreshed_items["KB-SECOND"]["decision_reason"], "保留另一条人工判断")

    def test_refresh_stale_item_excludes_source_removed_after_sync(self):
        clone = self._clone("clone-op-refresh-removed-source").get_json()
        batch_id = clone["review_batch_id"]
        item = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]
        with server.app.app_context():
            source = server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P20",
            ).one()
            server.db.session.delete(source)
            review_item = server.db.session.get(server.MatrixCloneReviewItem, item["item_id"])
            review_item.detection_status = "stale"
            review_item.apply_status = "stale"
            review_item.error_message = "源知识内容已变化，请重新扫描"
            server.db.session.commit()

        refreshed = self.client.post(
            f"/api/matrix/clone-review/items/{item['item_id']}/refresh",
        )

        self.assertEqual(refreshed.status_code, 200, refreshed.get_data(as_text=True))
        body = refreshed.get_json()
        self.assertTrue(body["success"])
        self.assertTrue(body["excluded"])
        self.assertEqual(body["item"]["detection_status"], "source_removed")
        self.assertFalse(body["item"]["requires_attention"])
        self.assertEqual(body["batch"]["expected_submit_changes"], [])

    def test_refresh_stale_batch_rechecks_existing_sources_and_excludes_removed_sources(self):
        with server.app.app_context():
            server.db.session.add(server.ProductMatrix(
                question_wiki_id="KB-SECOND",
                product_name="P20",
                is_configured=True,
                question_content="如何清洁尘盒？",
                answer_content="取出尘盒后清洁。",
                product_category="扫地机",
            ))
            server.db.session.commit()
        clone = self._clone("clone-op-refresh-stale-batch").get_json()
        batch_id = clone["review_batch_id"]
        with server.app.app_context():
            old_source = server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P20",
            ).one()
            old_source.answer_content = "最大吸力内容已同步更新。"
            removed_source = server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-SECOND", product_name="P20",
            ).one()
            server.db.session.delete(removed_source)
            review_items = server.MatrixCloneReviewItem.query.filter_by(batch_id=batch_id).all()
            for review_item in review_items:
                review_item.detection_status = "stale"
                review_item.apply_status = "stale"
                review_item.error_message = "源知识内容已变化，请重新扫描"
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            batch.error_message = "源知识内容已变化，请重新扫描后再完成或提交"
            server.db.session.commit()

        with patch.object(server, "load_ai_config", return_value={}):
            refreshed = self.client.post(f"/api/matrix/clone-review/{batch_id}/refresh-stale")

        self.assertEqual(refreshed.status_code, 200, refreshed.get_data(as_text=True))
        body = refreshed.get_json()
        self.assertEqual(body["refreshed_count"], 1)
        self.assertEqual(body["excluded_count"], 1)
        items = {item["question_wiki_id"]: item for item in body["items"]}
        self.assertEqual(items["KB-OLD"]["detection_status"], "scanned")
        self.assertEqual(items["KB-OLD"]["answer"], "最大吸力内容已同步更新。")
        self.assertEqual(items["KB-SECOND"]["detection_status"], "source_removed")
        self.assertFalse(items["KB-SECOND"]["requires_attention"])
        self.assertEqual(body["progress"]["scanned_count"], 1)

    def test_final_submit_rechecks_source_revision_before_writing_target_relation(self):
        clone = self._clone("clone-op-stale-submit").get_json()
        batch_id = clone["review_batch_id"]
        item = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]
        self.assertEqual(self.client.post(
            f"/api/matrix/clone-review/items/{item['item_id']}/decision",
            json={"decision": "keep"},
        ).status_code, 200)
        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={
            "manual_scope_confirmed": True, "rules_confirmed": True,
        })
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        with server.app.app_context():
            source = server.ProductMatrix.query.filter_by(question_wiki_id="KB-OLD", product_name="P20").one()
            source.answer_content = "最终提交前内容已更新。"
            server.db.session.commit()

        submitted = self.client.post("/api/matrix/submit_changes", json={
            "operation_id": "submit-stale-final",
            "changes": completed.get_json()["expected_submit_changes"],
        })

        self.assertEqual(submitted.status_code, 409, submitted.get_data(as_text=True))
        self.assertEqual(submitted.get_json()["code"], "knowledge_revision_stale")
        with server.app.app_context():
            self.assertIsNone(server.ProductMatrix.query.filter_by(
                question_wiki_id="KB-OLD", product_name="P30 Pro",
            ).first())
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            stale_item = server.MatrixCloneReviewItem.query.filter_by(batch_id=batch_id).one()
            self.assertEqual(batch.status, "reviewing")
            self.assertEqual(stale_item.detection_status, "stale")

    def test_parameter_refresh_failure_returns_reason_and_failed_progress(self):
        import requests

        operation_id = "clone-op-refresh-failed"
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None), patch.object(
            server, "_refresh_clone_parameter_comparison",
            side_effect=requests.RequestException("参数服务连接被拒绝"),
        ):
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": operation_id,
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True,
                    "auto_scan": True,
                    "detection_methods": ["parameter_check"],
                },
            })

        self.assertEqual(response.status_code, 503, response.get_data(as_text=True))
        body = response.get_json()
        self.assertEqual(body["code"], "parameter_refresh_unavailable")
        self.assertIn("参数服务连接被拒绝", body["detail"])
        progress = self.client.get(f"/api/matrix/clone_config/progress/{operation_id}").get_json()
        self.assertEqual(progress["status"], "failed")
        self.assertEqual(progress["phase"], "参数校对失败")
        self.assertEqual(progress["error_message"], "最新参数同步失败：参数服务连接被拒绝")

    def test_async_review_clone_returns_persisted_batch_before_generation(self):
        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None) as pull_source, patch.object(
            server, "_start_incremental_clone_review_batch"
        ) as start_worker:
            response = self.client.post("/api/matrix/clone_config", json={
                "operation_id": "clone-op-async-early-return",
                "mode": "model",
                "source": "P20",
                "targets": ["P30 Pro"],
                "strategy": "append",
                "scope": {"mode": "all"},
                "review": {
                    "enabled": True,
                    "async_processing": True,
                    "auto_scan": True,
                    "detection_methods": ["manual_difference_ai"],
                    "manual_difference_text": "旧型号最大吸力 22000Pa，新品 40000Pa",
                },
            })

        self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
        body = response.get_json()
        self.assertTrue(body["accepted"])
        self.assertEqual(body["review_status"], "generating")
        self.assertEqual(body["progress"]["generated_count"], 0)
        self.assertEqual(body["progress"]["generation_total"], 0)
        start_worker.assert_called_once()
        self.assertEqual(start_worker.call_args.args[0], body["review_batch_id"])
        pull_source.assert_not_called()
        detail = self.client.get(f"/api/matrix/clone-review/{body['review_batch_id']}").get_json()
        self.assertEqual(detail["status"], "generating")
        self.assertEqual(detail["items"], [])

        completed = self.client.post(f"/api/matrix/clone-review/{body['review_batch_id']}/complete", json={})
        applied = self.client.post(f"/api/matrix/clone-review/{body['review_batch_id']}/apply", json={})
        deleted = self.client.delete(f"/api/matrix/clone-review/{body['review_batch_id']}")
        self.assertEqual(completed.status_code, 409)
        self.assertEqual(completed.get_json()["code"], "clone_review_generation_incomplete")
        self.assertEqual(applied.status_code, 409)
        self.assertEqual(applied.get_json()["code"], "clone_review_generation_incomplete")
        self.assertEqual(deleted.status_code, 409)

    def test_rejects_decision_until_an_incremental_item_finishes_scanning(self):
        clone = self._clone("clone-op-pending-scan-decision").get_json()
        batch_id = clone["review_batch_id"]
        item_id = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]["item_id"]
        with server.app.app_context():
            item = server.db.session.get(server.MatrixCloneReviewItem, item_id)
            item.detection_status = "pending_scan"
            server.db.session.commit()

        response = self.client.post(
            f"/api/matrix/clone-review/items/{item_id}/decision",
            json={"decision": "keep"},
        )
        self.assertEqual(response.status_code, 409, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["code"], "clone_review_item_not_ready")

    def test_rejects_decision_while_ai_impact_review_is_pending(self):
        clone = self._clone("clone-op-ai-pending-decision").get_json()
        batch_id = clone["review_batch_id"]
        item_id = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]["item_id"]
        with server.app.app_context():
            item = server.db.session.get(server.MatrixCloneReviewItem, item_id)
            item.detection_status = "scanned"
            item.risk_level = "medium"
            item.decision = "pending"
            item.apply_status = "not_applied"
            item.ai_impact_review_json = '{"status":"processing","relation":"uncertain"}'
            server.db.session.commit()

        response = self.client.post(
            f"/api/matrix/clone-review/items/{item_id}/decision",
            json={"decision": "keep"},
        )
        self.assertEqual(response.status_code, 409, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["code"], "clone_review_ai_pending")

    def test_batch_payload_reports_persisted_generation_and_scan_counts(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="incremental-count-batch",
                operation_id="incremental-count-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                scope_snapshot_json='{"generation_total": 3}',
                detection_methods_json='["manual_difference_ai"]',
                status="generating",
                source_count=3,
                item_count=2,
                created_by="admin",
            )
            server.db.session.add(batch)
            server.db.session.add_all([
                server.MatrixCloneReviewItem(
                    item_id="incremental-pending",
                    batch_id=batch.batch_id,
                    question_wiki_id="KB-PENDING",
                    target_model="P30 Pro",
                    source_model="P20",
                    detection_status="pending_scan",
                    risk_level="unknown",
                    decision="pending",
                    apply_status="not_applied",
                ),
                server.MatrixCloneReviewItem(
                    item_id="incremental-scanned",
                    batch_id=batch.batch_id,
                    question_wiki_id="KB-SCANNED",
                    target_model="P30 Pro",
                    source_model="P20",
                    detection_status="scanned",
                    risk_level="default",
                    decision="keep",
                    apply_status="not_required",
                ),
            ])
            server.db.session.commit()

        detail = self.client.get("/api/matrix/clone-review/incremental-count-batch").get_json()
        self.assertEqual(detail["progress"]["generated_count"], 2)
        self.assertEqual(detail["progress"]["generation_total"], 3)
        self.assertEqual(detail["progress"]["scanned_count"], 1)
        self.assertFalse(detail["progress"]["generation_complete"])

    def test_resume_scan_only_processes_pending_scan_items(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="resume-pending-only-batch",
                operation_id="resume-pending-only-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json="[]",
                status="partial_failed",
                source_count=2,
                item_count=2,
                created_by="admin",
                error_message="模拟中断",
            )
            scanned = server.MatrixCloneReviewItem(
                item_id="resume-scanned-item",
                batch_id=batch.batch_id,
                question_wiki_id="KB-SCANNED-KEEP",
                target_model="P30 Pro",
                source_model="P20",
                question="已处理知识",
                answer="已保留",
                detection_status="scanned",
                risk_level="high",
                suggested_action="defer",
                decision="defer",
                decision_reason="保留人工判断",
                decided_by="admin",
                apply_status="not_applied",
            )
            pending = server.MatrixCloneReviewItem(
                item_id="resume-pending-item",
                batch_id=batch.batch_id,
                question_wiki_id="KB-PENDING-RESUME",
                target_model="P30 Pro",
                source_model="P20",
                question="待处理知识",
                answer="无需人工处理",
                detection_status="pending_scan",
                risk_level="unknown",
                suggested_action="defer",
                decision="pending",
                apply_status="not_applied",
            )
            server.db.session.add_all([batch, scanned, pending])
            server.db.session.commit()

        response = self.client.post(
            "/api/matrix/clone-review/resume-pending-only-batch/scan",
            json={"pending_only": True},
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        body = response.get_json()
        self.assertTrue(body["success"])
        self.assertEqual(body["progress"]["scanned_count"], 2)
        self.assertFalse(body["scan_active"])
        with server.app.app_context():
            scanned = server.db.session.get(server.MatrixCloneReviewItem, "resume-scanned-item")
            pending = server.db.session.get(server.MatrixCloneReviewItem, "resume-pending-item")
            self.assertEqual(scanned.risk_level, "high")
            self.assertEqual(scanned.decision, "defer")
            self.assertEqual(scanned.decision_reason, "保留人工判断")
            self.assertEqual(scanned.decided_by, "admin")
            self.assertEqual(pending.detection_status, "scanned")

    def test_resume_scan_rejects_batch_without_pending_items(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="resume-no-pending-batch",
                operation_id="resume-no-pending-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json="[]",
                status="reviewing",
                source_count=1,
                item_count=1,
                created_by="admin",
            )
            server.db.session.add(batch)
            server.db.session.add(server.MatrixCloneReviewItem(
                item_id="resume-no-pending-item",
                batch_id=batch.batch_id,
                question_wiki_id="KB-DONE",
                target_model="P30 Pro",
                source_model="P20",
                detection_status="scanned",
                risk_level="default",
                suggested_action="keep",
                decision="keep",
                apply_status="not_required",
            ))
            server.db.session.commit()

        response = self.client.post(
            "/api/matrix/clone-review/resume-no-pending-batch/scan",
            json={"pending_only": True},
        )
        self.assertEqual(response.status_code, 422, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["code"], "clone_review_no_pending_scan")

    def test_resume_scan_allows_pending_items_after_other_items_were_applied(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="resume-after-applied-batch",
                operation_id="resume-after-applied-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json="[]",
                status="partial_failed",
                source_count=2,
                item_count=2,
                created_by="admin",
            )
            applied = server.MatrixCloneReviewItem(
                item_id="resume-applied-item",
                batch_id=batch.batch_id,
                question_wiki_id="KB-APPLIED",
                target_model="P30 Pro",
                source_model="P20",
                detection_status="scanned",
                risk_level="high",
                suggested_action="keep",
                decision="keep",
                apply_status="applied",
            )
            pending = server.MatrixCloneReviewItem(
                item_id="resume-after-applied-pending",
                batch_id=batch.batch_id,
                question_wiki_id="KB-APPLIED-PENDING",
                target_model="P30 Pro",
                source_model="P20",
                question="待校验知识",
                answer="无需人工处理",
                detection_status="pending_scan",
                risk_level="unknown",
                suggested_action="defer",
                decision="pending",
                apply_status="not_applied",
            )
            server.db.session.add_all([batch, applied, pending])
            server.db.session.commit()

        response = self.client.post(
            "/api/matrix/clone-review/resume-after-applied-batch/scan",
            json={"pending_only": True},
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        body = response.get_json()
        self.assertTrue(body["success"])
        with server.app.app_context():
            self.assertEqual(
                server.db.session.get(server.MatrixCloneReviewItem, "resume-applied-item").apply_status,
                "applied",
            )
            self.assertEqual(
                server.db.session.get(server.MatrixCloneReviewItem, "resume-after-applied-pending").detection_status,
                "scanned",
            )

    def test_scan_does_not_overwrite_a_human_decision(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="human-decision-scan-batch",
                operation_id="human-decision-scan-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json='[]',
                status="scanning",
                source_count=1,
                item_count=1,
                created_by="admin",
            )
            item = server.MatrixCloneReviewItem(
                item_id="human-decision-scan-item",
                batch_id=batch.batch_id,
                question_wiki_id="KB-OLD",
                target_model="P30 Pro",
                source_model="P20",
                knowledge_revision=server._clone_review_revision("最大吸力是多少？", "最大吸力为 22000Pa。"),
                question="最大吸力是多少？",
                answer="最大吸力为 22000Pa。",
                product_category="扫地机",
                detection_status="pending_scan",
                risk_level="unknown",
                suggested_action="defer",
                decision="defer",
                decision_reason="人工已确认需要继续核对",
                decided_by="admin",
                apply_status="not_applied",
            )
            server.db.session.add_all([batch, item])
            server.db.session.commit()
            server._scan_clone_review_batch(batch, commit_every=1)
            server.db.session.commit()
            refreshed = server.db.session.get(server.MatrixCloneReviewItem, item.item_id)
            self.assertEqual(refreshed.decision, "defer")
            self.assertEqual(refreshed.decided_by, "admin")
            self.assertEqual(refreshed.decision_reason, "人工已确认需要继续核对")

    def test_generated_item_can_be_reviewed_while_batch_is_generating_but_cannot_finish_or_submit(self):
        clone = self._clone("clone-op-live-review").get_json()
        batch_id = clone["review_batch_id"]
        with server.app.app_context():
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            batch.status = "generating"
            batch.scope_snapshot_json = '{"generation_total": 2}'
            server.db.session.commit()

        detail = self.client.get(f"/api/matrix/clone-review/{batch_id}")
        self.assertEqual(detail.status_code, 200, detail.get_data(as_text=True))
        body = detail.get_json()
        self.assertEqual(body["status"], "generating")
        self.assertEqual(body["progress"]["generated_count"], 1)
        self.assertEqual(body["progress"]["generation_total"], 2)
        self.assertFalse(body["progress"]["generation_complete"])
        item = body["items"][0]

        decided = self.client.post(
            f"/api/matrix/clone-review/items/{item['item_id']}/decision",
            json={"decision": "keep", "decision_reason": "人工已核对当前生成内容"},
        )
        self.assertEqual(decided.status_code, 200, decided.get_data(as_text=True))
        self.assertEqual(decided.get_json()["item"]["decision"], "keep")

        completed = self.client.post(f"/api/matrix/clone-review/{batch_id}/complete", json={
            "manual_scope_confirmed": True,
            "rules_confirmed": True,
        })
        self.assertEqual(completed.status_code, 409, completed.get_data(as_text=True))
        self.assertEqual(completed.get_json()["code"], "clone_review_generation_incomplete")

        with patch.object(server, "_maybe_pull_matrix_and_logs_from_supabase", return_value=None):
            submitted = self.client.post("/api/matrix/submit_changes", json={
                "operation_id": "submit-live-review-too-early",
                "changes": [{
                    "question_wiki_id": item["question_wiki_id"],
                    "product_name": item["target_model"],
                    "old_is_configured": False,
                    "new_is_configured": True,
                    "edit_source": "bulk",
                }],
            })
        self.assertEqual(submitted.status_code, 409, submitted.get_data(as_text=True))
        self.assertEqual(submitted.get_json()["code"], "clone_review_required")

    def test_scan_preserves_human_decision_saved_while_processing(self):
        clone = self._clone("clone-op-preserve-live-decision").get_json()
        batch_id = clone["review_batch_id"]
        item_id = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]["item_id"]
        with server.app.app_context():
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            batch.status = "scanning"
            server.db.session.commit()

        decided = self.client.post(f"/api/matrix/clone-review/items/{item_id}/decision", json={
            "decision": "keep",
            "decision_reason": "人工判断可沿用",
        })
        self.assertEqual(decided.status_code, 200, decided.get_data(as_text=True))

        with server.app.app_context():
            batch = server.db.session.get(server.MatrixCloneReviewBatch, batch_id)
            server._scan_clone_review_batch(batch, commit_every=1)
            server.db.session.commit()

        item = self.client.get(f"/api/matrix/clone-review/{batch_id}").get_json()["items"][0]
        self.assertEqual(item["decision"], "keep")
        self.assertEqual(item["decision_reason"], "人工判断可沿用")
        with server.app.app_context():
            persisted = server.db.session.get(server.MatrixCloneReviewItem, item_id)
            self.assertEqual(persisted.decided_by, "admin")
            self.assertIsNotNone(persisted.decided_at)

    def test_incremental_scan_commits_each_item_without_holding_the_write_lock_for_a_chunk(self):
        with server.app.app_context():
            batch = server.MatrixCloneReviewBatch(
                batch_id="incremental-commit-batch",
                operation_id="incremental-commit-operation",
                source_mode="model",
                source_value="P20",
                target_models_json='["P30 Pro"]',
                detection_methods_json='[]',
                status="scanning",
                source_count=2,
                item_count=2,
                created_by="admin",
            )
            items = [
                server.MatrixCloneReviewItem(
                    item_id=f"incremental-commit-item-{index}",
                    batch_id=batch.batch_id,
                    question_wiki_id=f"KB-COMMIT-{index}",
                    target_model="P30 Pro",
                    source_model="P20",
                    question=f"测试知识 {index}",
                    answer="无需人工处理。",
                    detection_status="pending_scan",
                    decision="pending",
                    apply_status="not_applied",
                )
                for index in (1, 2)
            ]
            server.db.session.add_all([batch, *items])
            server.db.session.commit()

            observed_statuses = []
            original_set_progress = server._set_matrix_clone_progress

            def observe_committed_progress(operation_id, **kwargs):
                if kwargs.get("completed") == 2:
                    with sqlite3.connect(os.environ["KMATRIX_SQLITE_PATH"], timeout=0.1) as connection:
                        observed_statuses.append(connection.execute(
                            "SELECT detection_status FROM matrix_clone_review_item WHERE item_id = ?",
                            ("incremental-commit-item-1",),
                        ).fetchone()[0])
                return original_set_progress(operation_id, **kwargs)

            with patch.object(server, "_set_matrix_clone_progress", side_effect=observe_committed_progress):
                server._scan_clone_review_batch(batch, commit_every=25)

            self.assertEqual(observed_statuses, ["scanned"])


if __name__ == "__main__":
    unittest.main()
