#!/usr/bin/env python3
"""Export image link previews to an offline-labeling workbook.

The workbook keeps the source row id and URL unchanged so a returned workbook
can be matched back to link_previews without relying on row order.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import psycopg2
import requests
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from PIL import Image as PILImage


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CONFIG = ROOT / "KnowledgeBaseTool_Local" / "supabase_config_local.json"
DEFAULT_OUTPUT = Path(__file__).resolve().parent / "多媒体图片线下打标清单_20261016.xlsx"


def load_db_config(path: Path) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    config = payload.get("local_db") or {}
    if not config:
        raise RuntimeError(f"配置文件没有 local_db: {path}")
    return config


def fetch_rows(config: dict) -> list[dict]:
    query = """
        SELECT id, url, kb_id, type, tags, created_at
        FROM link_previews
        WHERE lower(COALESCE(type, '')) = 'image'
        ORDER BY id ASC
    """
    with psycopg2.connect(**config) as conn, conn.cursor() as cur:
        cur.execute(query)
        columns = [item.name for item in cur.description]
        return [dict(zip(columns, row)) for row in cur.fetchall()]


def safe_filename(url: str, index: int) -> str:
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".svg"}:
        suffix = ".img"
    return f"{index:04d}{suffix}"


def download_thumbnail(url: str, timeout_seconds: int = 30) -> tuple[bytes | None, str]:
    last_error = "请求失败"
    for attempt in range(3):
        try:
            response = requests.get(url, timeout=(8, timeout_seconds), headers={"User-Agent": "KMatrix-image-export/1.0"})
            response.raise_for_status()
            with PILImage.open(io.BytesIO(response.content)) as source:
                image = source.convert("RGB")
                image.thumbnail((320, 220), PILImage.Resampling.LANCZOS)
                output = io.BytesIO()
                image.save(output, format="JPEG", quality=86, optimize=True)
                return output.getvalue(), "成功"
        except Exception as exc:  # Keep the row even when one remote asset is unavailable.
            last_error = type(exc).__name__
            if attempt < 2:
                time.sleep(0.4 * (attempt + 1))
    return None, f"失败：{last_error}"


def clean_tags(value) -> str:
    if isinstance(value, list):
        return "、".join(str(item).strip() for item in value if str(item).strip())
    if value is None:
        return ""
    return str(value)


def split_tags(value) -> tuple[str, str]:
    """Separate system provenance tags from business tags for safe offline editing."""
    if isinstance(value, list):
        tags = [str(item).strip() for item in value if str(item).strip()]
    elif value is None:
        tags = []
    else:
        tags = [part.strip() for part in str(value).replace("，", ",").split(",") if part.strip()]
    system_tags = [tag for tag in tags if tag.startswith("来源:") or tag.startswith("外链类型:")]
    business_tags = [tag for tag in tags if tag not in system_tags]
    return "、".join(business_tags), "、".join(system_tags)


def format_created_at(value) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value)


def normalize_drawing_relationships(output: Path) -> None:
    """Use relative drawing targets so openpyxl and desktop Excel both reload images."""
    temporary = output.with_suffix(output.suffix + ".tmp")
    with ZipFile(output, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename.startswith("xl/drawings/_rels/") and entry.filename.endswith(".rels"):
                payload = payload.replace(b'Target="/xl/media/', b'Target="../media/')
            target.writestr(entry, payload)
    temporary.replace(output)


def anchor_image_to_cell(image: XLImage, row: int, column: int = 3) -> None:
    """Anchor the image to one worksheet cell so it moves and resizes with that cell."""
    image.anchor = TwoCellAnchor(
        editAs="twoCell",
        _from=AnchorMarker(col=column - 1, row=row - 1),
        to=AnchorMarker(col=column, row=row),
    )


def build_workbook(rows: list[dict], output: Path) -> tuple[int, int]:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "图片清单"
    instructions = workbook.create_sheet("填写说明")

    headers = [
        "序号", "记录ID（请勿修改）", "图片预览", "图片链接", "KB ID",
        "当前标签（可修改）", "来源标签（请勿修改）", "创建时间", "预览状态", "线下处理备注",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_blue = Side(style="thin", color="D9EAF7")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_blue)
    sheet.row_dimensions[1].height = 30

    # Remote image hosts occasionally stall. Parallel reads keep the export
    # bounded while each row still records its own preview status.
    with ThreadPoolExecutor(max_workers=16) as executor:
        previews = list(executor.map(download_thumbnail, [str(row.get("url") or "") for row in rows]))
    # A few CDN requests can time out only under concurrent load. Retry those
    # rows serially with a longer read timeout before leaving a missing preview.
    for index, (row, preview) in enumerate(zip(rows, previews)):
        if preview[0] is None:
            previews[index] = download_thumbnail(str(row.get("url") or ""), timeout_seconds=60)
    ok_count = 0
    failed_count = 0
    temporary_images: list[Path] = []
    for index, (row, preview) in enumerate(zip(rows, previews), start=1):
        image_data, preview_status = preview
        if image_data:
            ok_count += 1
        else:
            failed_count += 1
        business_tags, source_tags = split_tags(row.get("tags"))
        sheet.append([
            index,
            str(row.get("id") or ""),
            "",
            str(row.get("url") or ""),
            str(row.get("kb_id") or ""),
            business_tags,
            source_tags,
            format_created_at(row.get("created_at")),
            preview_status,
            "",
        ])
        current_row = sheet.max_row
        sheet.row_dimensions[current_row].height = 124
        sheet.cell(current_row, 4).hyperlink = str(row.get("url") or "")
        sheet.cell(current_row, 4).style = "Hyperlink"
        for column in (1, 2, 4, 5, 6, 7, 8, 9):
            sheet.cell(current_row, column).alignment = Alignment(vertical="top", wrap_text=True)
        for column in (9, 10):
            sheet.cell(current_row, column).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(current_row, 6).fill = PatternFill("solid", fgColor="FFF2CC")
        sheet.cell(current_row, 7).fill = PatternFill("solid", fgColor="F2F2F2")
        if image_data:
            image_path = output.parent / f".tmp_{safe_filename(str(row.get('url') or ''), index)}.jpg"
            image_path.write_bytes(image_data)
            temporary_images.append(image_path)
            embedded = XLImage(str(image_path))
            embedded.width = 250
            embedded.height = 172
            anchor_image_to_cell(embedded, current_row, column=3)
            sheet.add_image(embedded)

    widths = {1: 8, 2: 38, 3: 38, 4: 72, 5: 30, 6: 28, 7: 24, 8: 22, 9: 14, 10: 34}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.2, footer=0.2)
    sheet.print_title_rows = "1:1"

    instructions_rows = [
        ["用途", "图片记录线下打标，回传后按记录ID匹配回原数据。"],
        ["数据范围", f"当前运行库中 type=image 的图片记录，共 {len(rows)} 条。"],
        ["请勿修改", "记录ID（请勿修改）、图片链接、KB ID、来源标签（请勿修改）。"],
        ["请填写", "只修改当前标签（可修改）：填写最终业务标签，多个标签用顿号、逗号或换行分隔。"],
        ["请填写", "线下处理备注：记录需要删除、替换、确认或其他处理意见。"],
        ["图片预览", "工作簿内嵌的是缩略图；原图请点击图片链接打开。预览状态失败的记录仍保留。"],
        ["回传建议", "请保留 Excel 文件中的记录ID和原始列，不要排序后删除记录；直接回传填写后的文件即可。"],
        ["导出时间", datetime.now().astimezone().isoformat(timespec="seconds")],
        ["预览统计", f"成功 {ok_count} 条，失败 {failed_count} 条"],
    ]
    for row in instructions_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 110
    for row in instructions.iter_rows():
        row[0].font = Font(bold=True, color="1F4E78")
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
        instructions.row_dimensions[row[0].row].height = 32
    instructions.sheet_view.showGridLines = False
    instructions.freeze_panes = "A2"

    workbook.properties.title = "多媒体图片线下打标清单"
    workbook.properties.subject = "全部链接中的图片记录"
    workbook.properties.creator = "KnowBase Hub"
    workbook.save(output)
    normalize_drawing_relationships(output)
    for image_path in temporary_images:
        image_path.unlink(missing_ok=True)
    return ok_count, failed_count


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=Path(os.environ.get("KMATRIX_CONFIG_DIR", "")) / "supabase_config_local.json")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    config_path = args.config if args.config.exists() else DEFAULT_CONFIG
    rows = fetch_rows(load_db_config(config_path))
    success, failed = build_workbook(rows, args.output)
    print(json.dumps({"output": str(args.output), "rows": len(rows), "preview_ok": success, "preview_failed": failed}, ensure_ascii=False))


if __name__ == "__main__":
    main()
